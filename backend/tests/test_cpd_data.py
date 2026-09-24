from datetime import datetime
from types import SimpleNamespace
import stat
from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel
import pytest
from apps.reports.cpd_data_service import data_root, safe_file, walk_xlsx, inspect_sheet, preview_sheet, time_window

def workbook():
    book=Workbook(); sheet=book.active; sheet.title='Results'
    sheet.append(['Time','value'])
    for ts,n in [('2026-09-23 10:00:00',3),('2026-09-23 10:00:01',1),('2026-09-23 10:00:02',2),('2026-09-23 10:00:03',4)]: sheet.append([ts,n])
    sheet.append(['invalid',0])
    return book

def test_root_includes_remote_username_and_prevents_escape():
    root=data_root(SimpleNamespace(upper_machine=SimpleNamespace(username='tester')),'CPD','Temperature')
    assert root=='/data/tester/cpd_data/cpd/temperature'
    assert safe_file(root,root+'/deep/result.xlsx').endswith('result.xlsx')
    with pytest.raises(ValueError): safe_file(root,root+'/../other.xlsx')

def test_inclusive_internal_time_and_numeric_sort_pagination():
    book=workbook();start,end=time_window('2026-09-23 10:00:00','2026-09-23 10:00:02')
    preview=preview_sheet(book,'Results',start,end,page=2,page_size=2,sort_column=1)
    assert preview['total']==3 and preview['rows'][0]['cells'][1]==3
    assert preview['rows'][0]['line']==2
    assert preview['invalid_time_rows']==1

def test_excel_serial_and_missing_time_column():
    book=Workbook();s=book.active;s.append(['Timestamp','n']);s.append([to_excel(datetime(2026,9,23,10,0)),4])
    start,end=time_window('2026-09-23 10:00:00','2026-09-23 10:00:00')
    assert inspect_sheet(s,book.epoch,start,end)['matched_rows']==1
    missing=book.create_sheet('Summary');missing.append(['temperature',99])
    assert preview_sheet(book,'Summary',start,end)['filter_status']=='missing_time_column'

def test_recursive_scan_ignores_mtime_and_links():
    def entry(name,mode): return SimpleNamespace(filename=name,st_mode=mode,st_size=100,st_mtime=0)
    folders={'/r':[entry('one',stat.S_IFDIR)],'/r/one':[entry('two',stat.S_IFDIR)],'/r/one/two':[entry('x.xlsx',stat.S_IFREG),entry('outside',stat.S_IFLNK)]}
    result=list(walk_xlsx(SimpleNamespace(listdir_attr=lambda p:folders[p]),'/r'))
    assert [r['path'] for r in result]==['/r/one/two/x.xlsx']
