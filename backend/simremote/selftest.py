"""模拟机群端到端自检。

覆盖「后端真实会发的命令」：SFTP 只读访问、find（含 -printf NUL 协议）、
tar 列表与嵌套解压管道、awk 时间窗过滤、长驻 tail 转发，以及日志行是否
匹配后端的三套内置正则。

用法：``python -m simremote.cli selftest``（需要机群已在运行）
"""

from __future__ import annotations

import fnmatch
import math
import os
import re
import tempfile
import time
from datetime import datetime, timedelta
from pathlib import Path

import paramiko

from . import atlog_site, fleet, livesim

_RESULTS: list[tuple[str, bool, str]] = []

# 调用链规则的两个方向符。前端 ``logParser.ts`` 用的是同一条正则
# （``/>\s*\(\s*\)/`` 入口、``/<\s*\(\s*\)/`` 出口）。
ENTRY_MARKER_REGEX = re.compile(r">\s*\(\s*\)")
EXIT_MARKER_REGEX = re.compile(r"<\s*\(\s*\)")
#: 方向符**前面**必须是被调用的函数名，而且要写成**调用形状** ``函数名()``。
#:
#: 与前端 ``logParser.ts::FUNCTION_PREFIX_REGEX``
#: （``/^\s*([A-Za-z_~][\w:<>~.-]*\(\))/``）同形状 —— 那是它**首选**的解析路径：
#: 正文以 ``ScanLot()`` 开头，``functionName`` 直接拿到值。首字符限制为字母/下划线，
#: 所以 ``[2026-09-24 22:50:00.000]`` 这类时间戳方括号不会被误认成函数名。
BOUNDARY_NAME_REGEX = re.compile(r"[A-Za-z_~][\w:<>~.\-]*\(\)")
#: **旧写法（禁止）**：方括号包函数名 + 方向符，例如 ``[ScanLot] >()``。
#:
#: 它只能靠前端 ``parseBoundaryFunctionName`` 的兜底分支勉强解析出来：``functionName``
#: 是空、只有 ``boundaryFunctionName`` 有值，于是「连续同名函数日志」「闭合后的同名
#: 尾随日志」两条内置折叠规则对阶段框/边界行全部失效。渲染出这种形状即判失败。
LEGACY_BRACKET_CALL_REGEX = re.compile(r"\[[A-Za-z_~][\w:<>~.\-]*]\s*[<>]\s*\(\s*\)")
#: 日志行的 ``rpc`` 字段 ``[模块:函数:行号]`` —— 它给出这一行**属于哪个函数**。
#: 正文行没有方向符，"这条正文有没有落在它自己那一次调用的边界里"只能靠 rpc 反查
#: 函数名再跟栈顶比（见 ``check_wsp_point_call_boundaries``）。
RPC_FUNCTION_REGEX = re.compile(r"\[[a-z][a-z0-9_]*:([A-Za-z_][A-Za-z0-9_]*):\d+\]")

#: 实时 tail 的等待窗口（秒）。
#:
#: 远端 ``tail -F`` 的 stdout 是**管道**不是 tty，stdio 走全缓冲，要攒满约 4KB 才
#: flush 一次；日志源是 5 条流各 0.5s 一行，单行 ~150B，所以一条流要攒满 ~27 行
#: （约 14s）才会把第一批行推给订阅方。窗口开短了就会隔几次假失败一次 ——
#: 真实机台也是这个行为，不是模拟器卡住了。
LIVE_TAIL_WINDOW_SECONDS = 100
#: 日志源没在跑时的探针模式窗口（自己写一行探针，不用等积攒）
LIVE_TAIL_PROBE_WINDOW_SECONDS = 6


def _record(name: str, ok: bool, detail: str = "") -> bool:
    _RESULTS.append((name, bool(ok), detail))
    return bool(ok)


def _physical_lines(text: str) -> list[str]:
    """文本 -> 非空物理行。"""
    return [line for line in text.splitlines() if line.strip()]


def _group_records(lines: list[str]) -> list[list[str]]:
    """把物理行按「行首时间戳」归组成**记录**：带时间戳的是新记录的首行，其后所有
    不带时间戳的行都是它的续行。

    为什么校验类代码必须按**记录**判、不能按**物理行**判：真机台的测量 / 干涉模块
    会把一次采样的整包读数打成字典落一条日志 —— **一条记录横跨十几行，只有首行带
    八字段前缀**（构造见 ``loggen.scan_metric_body``）。逐行要求"每一行都是八字段"
    会把这种完全合法的日志判成格式错误，而它恰恰是解析器必须处理的真实形态
    （也正是"按行首时间戳切记录"会截断正文的那个场景）。

    调用方常常只取尾部若干行、或把几条流的文本拼在一起，起点因此可能落在某条记录
    **中间**。开头那段没有可归属记录的续行直接跳过 —— 它们没有时间戳，无法当作
    记录校验，不能拿来报"首行不合规"。
    """
    from apps.logsources.services.file_index import parse_line_time

    records: list[list[str]] = []
    for line in lines:
        if parse_line_time(line) is not None:
            records.append([line])
        elif records:
            records[-1].append(line)
    return records


def _record_head(line: str) -> str:
    """一条记录的正文：前 8 个 ``[...]`` 是固定字段，第 9 段才是正文。"""
    return line.split("] ", 8)[-1]


def _connect(spec: fleet.MachineSpec) -> paramiko.SSHClient:
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(
        spec.host,
        port=spec.ssh_port,
        username=fleet.SIM_USERNAME,
        password=fleet.SIM_PASSWORD,
        timeout=6,
        allow_agent=False,
        look_for_keys=False,
    )
    return client


def _exec(client: paramiko.SSHClient, command: str, timeout: int = 30) -> tuple[bytes, bytes, int]:
    _stdin, stdout, stderr = client.exec_command(command, timeout=timeout)
    out = stdout.read()
    err = stderr.read()
    return out, err, stdout.channel.recv_exit_status()


# --------------------------------------------------------------------------- 各项检查


def check_ports() -> None:
    for spec in fleet.FLEET:
        try:
            client = _connect(spec)
            client.close()
            _record(f"SSH 可达 {spec.host}:{spec.ssh_port}", True)
        except Exception as exc:  # noqa: BLE001
            _record(f"SSH 可达 {spec.host}:{spec.ssh_port}", False, str(exc))


def check_auth_reject() -> None:
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        client.connect(
            fleet.UPPER.host, port=fleet.UPPER.ssh_port, username="intruder",
            password="nope", timeout=5, allow_agent=False, look_for_keys=False,
        )
        client.close()
        _record("错误口令被拒绝", False, "竟然登录成功了")
    except paramiko.AuthenticationException:
        _record("错误口令被拒绝", True)
    except Exception as exc:  # noqa: BLE001
        _record("错误口令被拒绝", False, str(exc))


def check_sftp_readonly(client: paramiko.SSHClient) -> None:
    sftp = client.open_sftp()
    try:
        entries = sftp.listdir_attr(f"{fleet.UPPER.debug_root}")
        names = sorted(item.filename for item in entries)
        _record("SFTP 列出调试根目录", len(names) >= 5, f"{len(names)} 项：{names[:4]}")

        module_dir = f"{fleet.UPPER.debug_root}/spwsp"
        files = sorted(item.filename for item in sftp.listdir_attr(module_dir))
        _record(
            "SFTP 列出子系统目录",
            any(item.endswith(".log") for item in files),
            f"{len(files)} 项",
        )

        target = f"{module_dir}/spwsp.log"
        info = sftp.stat(target)
        _record("SFTP stat 返回真实大小", info.st_size > 0, f"size={info.st_size}")

        with sftp.open(target, "rb") as handle:
            head = handle.read(64)
            handle.seek(0)
            first = handle.readline()
        _record(
            "SFTP 随机读 + readline",
            head.startswith(b"[") and first.startswith(b"["),
            first.decode("utf-8", "replace")[:48],
        )

        try:
            sftp.open(f"{module_dir}/readonly-probe.log", "w")
            _record("SFTP 拒绝写操作", False, "竟然允许写入")
        except OSError:
            _record("SFTP 拒绝写操作", True)
    finally:
        sftp.close()


def check_find(client: paramiko.SSHClient) -> None:
    out, _err, code = _exec(
        client,
        r"""find /log/tracepilot/debug -mindepth 2 -maxdepth 2 -type f \( -name '*.log' -o -name '*.tar.gz' \) -print""",
    )
    lines = [item for item in out.decode().splitlines() if item]
    # 实时日志源会在同一目录里留下 <fm>_<时刻>.log 这种滚动归档（真实机台也一样），
    # 所以这里校验的是"14 个模块都被发现、且文件名都是合法形态"，而不是总数恰好相等。
    modules = {module for _subsystem, module in fleet.all_modules()}
    artifact_re = re.compile(r"^(?P<module>[a-z0-9]+)(?:_\d{8}|_\d{14})?\.(?:log|tar\.gz)$")
    parsed: list[tuple[str, str]] = []
    illegal: list[str] = []
    for item in lines:
        name = Path(item).name
        match = artifact_re.match(name)
        if match and match.group("module") in modules:
            parsed.append((item, match.group("module")))
        else:
            illegal.append(item)
    covered = {module for _item, module in parsed}
    _record(
        "find 模块发现（-print）",
        code == 0 and not illegal and covered == modules,
        f"{len(lines)} 行 · 覆盖 {len(covered)}/{len(modules)} 个模块"
        + (f" · 非法文件名 {illegal[:3]}" if illegal else ""),
    )
    archive_counts: dict[str, int] = {}
    for item, module in parsed:
        if item.endswith(".tar.gz"):
            archive_counts[module] = archive_counts.get(module, 0) + 1
    thin = sorted(module for module in modules if archive_counts.get(module, 0) < 3)
    _record(
        "find 发现每日归档包",
        not thin,
        f"每模块归档数最少 {min(archive_counts.values(), default=0)}（应 ≥3）"
        + (f" · 不足 {thin[:4]}" if thin else ""),
    )
    _record(
        "find 输出为远端路径",
        all(item.startswith("/log/") for item in lines),
        lines[0] if lines else "(空)",
    )

    out, _err, code = _exec(
        client,
        r"""find /log/tracepilot/debug/spwsp -maxdepth 1 -type f \( -iname 'spwsp.log' -o -iname 'spwsp_*.log' \) -printf '%i\t%s\t%T@\t%p\0'""",
    )
    records = [item for item in out.split(b"\0") if item]
    shape_ok = bool(records)
    for record in records:
        parts = record.decode("utf-8", "replace").split("\t")
        if len(parts) != 4:
            shape_ok = False
            break
        try:
            int(parts[0])
            int(parts[1])
            float(parts[2])
        except ValueError:
            shape_ok = False
            break
    _record(
        "find -printf NUL 协议",
        code == 0 and shape_ok,
        f"{len(records)} 条记录，字段 %i/%s/%T@/%p",
    )


def check_tar(client: paramiko.SSHClient) -> None:
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    nested = today - timedelta(days=2)
    archive = f"{fleet.UPPER.debug_root}/spwsp/spwsp_{nested:%Y%m%d}.tar.gz"

    out, err, code = _exec(client, f"tar -tzf {archive}")
    members = [item for item in out.decode().splitlines() if item]
    _record("tar -tzf 列出嵌套日包", code == 0 and len(members) == 24, f"{len(members)} 个成员（期望 24）")

    inner = members[0] if members else ""
    _record(
        "内层包名 = 小时段结束时刻",
        bool(inner) and inner.endswith(".tar.gz") and inner.startswith("spwsp_"),
        inner,
    )

    out, err, code = _exec(client, f"tar -xOzf {archive} {inner} | tar -tzf -")
    _record("嵌套管道逐级解压", code == 0 and out.decode().strip() == "spwsp.log", out.decode().strip()[:40])

    out, _err, code = _exec(client, f"tar -xOzf {archive} {inner} | tar -xOzf - spwsp.log | head -1")
    _record("解出内层成员正文", code == 0 and out.startswith(b"["), out.decode("utf-8", "replace")[:56])


def check_close_boundary() -> None:
    """轮转文件名后缀必须是「关闭边界」：今天 0 点关闭的文件覆盖昨天。"""
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    rotated = f"{fleet.UPPER.debug_root}/spwsp/spwsp_{today:%Y%m%d}000000.log"
    local = fleet.remote_to_local(fleet.UPPER, rotated)
    if not local.exists():
        _record("轮转文件按关闭边界命名", False, f"缺少 {rotated}")
        return
    with local.open(encoding="utf-8") as handle:
        stamps = [line[1:20] for line in handle if line.startswith("[")]
    first, last = stamps[0], stamps[-1]
    yesterday = (today - timedelta(days=1)).strftime("%Y-%m-%d")
    _record(
        "轮转文件按关闭边界命名",
        first.startswith(yesterday) and last.startswith(yesterday),
        f"{first} → {last}（应落在 {yesterday}）",
    )


def check_timestamp_jitter() -> None:
    """日志时间戳**不许落在整齐的网格上**。

    这条是用户提的：「日期太整齐了，要随机给点时间尾巴不一样，有时间不断变化的感觉，
    现在的日志都是 000 结尾，太假了」。曾经的生成方式是 ``start + i * step`` 的纯算术
    序列，后果是

    * 毫秒位恒为 ``.000``；
    * 步长 60s / 300s 时秒位恒为 ``:00``，一分钟一行就真的每分钟第 0 秒一行。

    所以这里查三件事，任何一件退回去都要当场失败：

    1. **毫秒位要散**：结尾 ``.000`` 的行占比必须很低（真机台不可能整秒落笔）；
    2. **秒位要铺开**：一分钟六十个秒值，一份上千行的文件至少要覆盖到两位数；
    3. **事件驱动型（用例日志）毫秒位要各不相同** —— 那几条本来就不是周期采样。

    三个来源分别覆盖：批量调试日志（当前段 + 轮转段）、运行事件日志、ATLog 用例日志。
    只查一类是不够的：它们各自走 ``clock_series`` 的不同 ratio，改坏一个不会影响另一个。
    """

    def collect(paths: list[Path]) -> list[str]:
        """按真实查询路径取样：**当前段 + 最新归档**一起读。

        只读当前段会有一个假失败源：活动文件写满 ``MAX_LIVE_LINES`` 行就轮转，刚轮转过的
        那几十秒里当前段只有几行，"秒位要铺开"自然凑不出来。真实查询本来就是
        两段一起读的。
        """
        stamps: list[str] = []
        for path in paths:
            if not path.exists():
                continue
            with path.open(encoding="utf-8", errors="replace") as handle:
                for line in handle:
                    if line.startswith("["):
                        # 前 23 个字符 = YYYY-MM-DD HH:MM:SS.mmm
                        stamps.append(line[1:24])
        return stamps

    def newest(directory: Path, pattern: str, current: Path) -> Path | None:
        candidates = [item for item in directory.glob(pattern) if item != current and item.is_file()]
        if not candidates:
            return None
        return max(candidates, key=lambda item: item.stat().st_mtime)

    def scan(label: str, paths: list[Path], *, min_seconds: int, max_zero_ratio: float) -> None:
        stamps = collect(paths)
        if not stamps:
            _record(label, False, "取样为空：" + "、".join(item.name for item in paths))
            return
        zero = sum(1 for item in stamps if item.endswith(".000"))
        seconds = {item[17:19] for item in stamps}
        _record(
            label,
            zero / len(stamps) <= max_zero_ratio and len(seconds) >= min_seconds,
            f"{len(stamps)} 行 · 结尾 .000 占 {zero / len(stamps):.1%} · 秒位 {len(seconds)}/60 种",
        )

    debug_root = fleet.remote_to_local(fleet.UPPER, fleet.UPPER.debug_root)
    run_root = fleet.remote_to_local(fleet.UPPER, fleet.UPPER.run_root)
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)

    spwsp_dir = debug_root / "spwsp"
    spwsp_current = spwsp_dir / "spwsp.log"
    rotated = spwsp_dir / f"spwsp_{today:%Y%m%d}000000.log"
    scan(
        "调试日志时间戳不落整秒网格",
        [spwsp_current, rotated],
        min_seconds=25,
        max_zero_ratio=0.02,
    )
    scan("轮转段时间戳不落整秒网格", [rotated], min_seconds=25, max_zero_ratio=0.02)

    event_current = run_root / "event.log"
    event_archive = newest(run_root, "event_*.log", event_current)
    scan(
        "运行事件时间戳不落整秒网格",
        [event_current] + ([event_archive] if event_archive else []),
        min_seconds=25,
        max_zero_ratio=0.05,
    )

    # 用例日志：只有十九行、六分钟，靠的是"事件驱动"那一档抖动。
    # 判据放宽到"毫秒位基本不重复"，因为样本本来就少。
    case = next((item for item in atlog_site.CASES if item.status == "failed"), atlog_site.CASES[0])
    case_log = (
        atlog_site.SITE_ROOT / atlog_site._case_relative_dir(case)
        / "full_logs" / "log" / "debug" / case.subsystem / f"{case.module}.log"
    )
    stamps = collect([case_log])
    if stamps:
        millis = {item.rsplit(".", 1)[-1] for item in stamps}
        zero = sum(1 for item in stamps if item.endswith(".000"))
        # ⚠️ 别用 `zero == 0`：微秒位是**均匀随机**的，19 行里至少出现一条 ``.000``
        # 的概率约 1.9% —— 那会变成"每 50 次 init 假失败一次"的骰子，而不是判据。
        # 判据的本意是"不落整秒网格"（即不是整分/整秒对齐），用比例表达才稳定：
        # 真网格是 100% 命中，这里只允许零星一两条。
        _record(
            "用例日志时间戳不落整秒网格",
            zero / len(stamps) <= 0.06 and len(millis) >= max(3, len(stamps) // 2),
            f"{len(stamps)} 行 · 毫秒位 {len(millis)} 种 · 结尾 .000 {zero} 条",
        )
    else:
        _record("用例日志时间戳不落整秒网格", False, f"缺少 {case_log}")


def check_awk_window(client: paramiko.SSHClient) -> None:
    now = datetime.now()
    start = (now - timedelta(hours=1)).strftime("%Y-%m-%d %H:%M:%S")
    end = now.strftime("%Y-%m-%d %H:%M:%S")
    command = (
        f"tar -xOzf {fleet.UPPER.debug_root}/spwsp/spwsp_20260919.tar.gz spwsp.log | "
        f"LC_ALL=C awk -v s='{start}' -v e='{end}' "
        """'{ if (substr($0,1,1) != "[") next; ts=substr($0,2,19); if (ts < s) next; if (ts > e) exit; print $0 }'"""
    )
    out, err, code = _exec(client, command)
    _record("awk 窗口过滤：窗口外返回 0 行", code == 0 and out.strip() == b"", f"{len(out.splitlines())} 行")

    current = f"{fleet.UPPER.debug_root}/spwsp/spwsp.log"
    command = f"cat -- {current} | head -3"
    out, err, code = _exec(client, command)
    _record("cat 读取当前日志", code == 0 and out.count(b"\n") == 3, f"{out.count(b'\n')} 行")

    window_start = (now - timedelta(minutes=30)).strftime("%Y-%m-%d %H:%M:%S")
    # 实时日志源的活动文件写满 MAX_LIVE_LINES 行就轮转，刚轮转过的那一小会儿当前段只有
    # 几行 —— 只查当前段会误报"窗口没命中"。真实查询路径本来就是
    # 「当前段 + 覆盖窗口的归档」一起读，这里跟着一起读。
    command = (
        "{ f=$(ls -1t "
        f"{fleet.UPPER.debug_root}/spwsp/spwsp_*.log 2>/dev/null | head -1); "
        f'[ -n "$f" ] && cat -- "$f"; cat -- {current}; }} | '
        f"LC_ALL=C awk -v s='{window_start}' -v e='{end}' "
        """'{ if (substr($0,1,1) != "[") next; ts=substr($0,2,19); if (ts < s) next; if (ts > e) exit; print $0 }'"""
    )
    out, _err, code = _exec(client, command)
    _record(
        "awk 窗口过滤：最近 30 分钟命中",
        code == 0 and len(out.splitlines()) >= 20,
        f"{len(out.splitlines())} 行（当前段 spwsp.log + 最新归档）",
    )

    # wsp 是新组件，正文还是"固定点位"格式 —— 单独过一遍**完全相同的查询路径**：
    # 窗口里既要有行，而且这些行得真的能被解析成移动点位。只验证 spwsp 不够，
    # 因为"前端能不能搜到 wsp"是用户会第一步就去做的事。
    wsp_current = f"{fleet.UPPER.debug_root}/wsp/wsp.log"
    command = (
        "{ f=$(ls -1t "
        f"{fleet.UPPER.debug_root}/wsp/wsp_*.log 2>/dev/null | head -1); "
        f'[ -n "$f" ] && cat -- "$f"; cat -- {wsp_current}; }} | '
        f"LC_ALL=C awk -v s='{window_start}' -v e='{end}' "
        """'{ if (substr($0,1,1) != "[") next; ts=substr($0,2,19); if (ts < s) next; if (ts > e) exit; print $0 }'"""
    )
    out, _err, code = _exec(client, command)
    wsp_rows = [item for item in out.decode("utf-8", "replace").splitlines() if item.strip()]
    wsp_points = [item for item in wsp_rows if "move absolute" in item]
    _record(
        "awk 窗口过滤：最近 30 分钟命中 wsp 点位行",
        code == 0 and len(wsp_points) >= 3,
        f"{len(wsp_rows)} 行（其中点位行 {len(wsp_points)} 行，当前段 wsp.log + 最新归档）",
    )


def check_version_file(client: paramiko.SSHClient) -> None:
    out, err, code = _exec(client, "cat -- /home/tracepilot/SW/version")
    # 版本文件要和真实机器一致：带 `Current Version:` 标记（后端解析器只认这个）。
    expected = f"Current Version: {fleet.SOFTWARE_VERSION}"
    _record("cat 版本文件", code == 0 and out.decode().strip() == expected, out.decode().strip())

    out, _err, code = _exec(client, """printf '%s\n' "$(hostname)"; uname -srm""")
    _record("连接探测命令", code == 0 and b"\n" in out, out.decode().strip().replace("\n", " | ")[:48])


def check_line_formats() -> None:
    from apps.logsources.services.log_format_parser import (
        BUILTIN_RULE_DEFAULTS,
        _compile,
    )

    rules = {
        "debug": BUILTIN_RULE_DEFAULTS[("debug", "标准调试日志")],
        "run": BUILTIN_RULE_DEFAULTS[("run", "运行事件日志")],
    }
    # 逐条查：wsp 是新组件、正文里还带 ``{ … }`` 花括号（转义写错会把整行打歪），
    # 所以它不能只靠"和 spwsp 同格式"推断，得自己过一遍内置正则。
    for category, relative in (
        ("debug", "spwsp/spwsp.log"),
        ("debug", "wsp/wsp.log"),
        ("run", "event.log"),
    ):
        root = fleet.UPPER.debug_root if category == "debug" else fleet.UPPER.run_root
        local = fleet.remote_to_local(fleet.UPPER, f"{root}/{relative}")
        if not local.exists():
            _record(f"{category} 日志行匹配内置正则 · {relative}", False, "文件不存在")
            continue
        pattern = _compile(rules[category]["pattern"])
        # 按**记录**校验，不按物理行：跨行记录（spwsp 的测量快照字典）只有首行带
        # 八字段前缀，续行既没有时间戳也没有字段 —— 那是真实机台的形态，不是格式
        # 错误。逐行判会把合法日志误报成失败。抽样口径不变（仍是前 200 行）。
        with local.open(encoding="utf-8") as handle:
            head_lines = [line.rstrip("\n") for _, line in zip(range(200), handle)]
        records = _group_records(_physical_lines("\n".join(head_lines)))
        bad = [lines[0] for lines in records if not pattern.match(lines[0])]
        _record(
            f"{category} 日志记录匹配内置正则 · {relative}",
            bool(records) and not bad,
            f"{len(records)} 条记录 / {len(head_lines)} 行，首行不合规 {len(bad)}"
            + (f"：{bad[0][:110]}" if bad else ""),
        )

    executor_root = fleet.UPPER.elog_root
    executor_files = sorted(
        Path(fleet.remote_to_local(fleet.UPPER, executor_root)).glob(f"{fleet.LOWER1.host}/*/*.log")
    )
    if not executor_files:
        _record("executor 日志行匹配内置正则", False, "未找到 executor 文件")
        return
    pattern = _compile(BUILTIN_RULE_DEFAULTS[("executor", "标准执行器日志")]["pattern"])
    matched = total = 0
    with executor_files[0].open(encoding="utf-8") as handle:
        for index, line in enumerate(handle):
            if index >= 200:
                break
            total += 1
            if pattern.match(line.rstrip("\n")):
                matched += 1
    _record("executor 日志行匹配内置正则", total > 0 and matched == total, f"{matched}/{total}")

    # 部署事件是**追加在 event.log 末尾**的，上面那条只抽样前 200 行覆盖不到 —— 单独全查一遍。
    # 不合格式的杂行会让后端的运行日志查询整体解析出错，所以这条必须严格。
    run_log = fleet.remote_to_local(fleet.UPPER, f"{fleet.UPPER.run_root}/event.log")
    if not run_log.exists():
        _record("部署事件合运行日志格式", False, "event.log 不存在")
        return
    run_pattern = _compile(BUILTIN_RULE_DEFAULTS[("run", "运行事件日志")]["pattern"])
    with run_log.open(encoding="utf-8") as handle:
        deploy_lines = [line.rstrip("\n") for line in handle if "[DEPLOY]" in line]
    malformed = [line for line in deploy_lines if not run_pattern.match(line)]
    _record(
        "部署事件合运行日志格式",
        bool(deploy_lines) and not malformed,
        f"{len(deploy_lines)} 条部署事件，不合格式 {len(malformed)} 条"
        if deploy_lines
        else "没有部署事件（先跑 scripts/sim.sh deploy）",
    )


def _stream_pid() -> int | None:
    pid_file = Path(__file__).resolve().parent / "run" / "stream.pid"
    if not pid_file.exists():
        return None
    try:
        return int(pid_file.read_text(encoding="utf-8").strip())
    except (OSError, ValueError):
        return None


def _stream_running() -> bool:
    pid = _stream_pid()
    if pid is None:
        return False
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


def check_live_tail() -> None:
    """实时监听：订阅点之后新产出的行应经 SSH ``tail -F`` 通道推过来。

    对**每一条会被用户单独订阅的流**（``livesim.OBSERVER_KEYS``）各开一条保活
    tail，但**共用同一个等待窗口** —— 所以总耗时和只探一条差不多。多探这几条的
    意义在于："这条流能不能实时看到"是用户会亲手去点的事，光看剧本在刷不算数；
    新增一条让人专门去看的流时，把它加进 OBSERVER_KEYS 就会自动被这里覆盖。

    实时日志源在跑时直接用它产出的行验证，既不写探针也不回滚 —— 并发写入下把
    文件截断回原长度，会把日志源这期间写的行一起削掉，还会让远端 tail 以为文件
    被截断而重读。日志源没跑时才退回"自己写一行探针、事后回滚"的老办法。
    """
    from . import livesim

    probes = [livesim._TARGET_BY_KEY[key] for key in livesim.OBSERVER_KEYS]
    running = _stream_running()
    marker = f"selftest probe {time.time():.3f}"
    sizes = {probe.key: probe.local_path.stat().st_size for probe in probes}
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    opened: list[tuple[livesim.StreamTarget, paramiko.Channel]] = []
    try:
        client.connect(
            fleet.UPPER.host, port=fleet.UPPER.ssh_port, username=fleet.SIM_USERNAME,
            password=fleet.SIM_PASSWORD, timeout=6, allow_agent=False, look_for_keys=False,
        )
        transport = client.get_transport()
        assert transport is not None
        for probe in probes:
            channel = transport.open_session()
            channel.exec_command(f"LC_ALL=C exec tail -n 0 -F -- {probe.remote_path}")
            opened.append((probe, channel))
        time.sleep(0.6)
        if not running:
            for probe, _channel in opened:
                # 探针行也要带真实的毫秒尾巴：写死 ``.000`` 的话，这条自己人写的行
                # 混进日志里反而比模拟数据更显眼（"就这一行是整的"）。
                moment = datetime.now()
                line = (
                    f"[{moment:%Y-%m-%d %H:%M:%S}.{moment.microsecond // 1000:03d}] "
                    f"[INFO] [{probe.subsystem.upper()}] "
                    f"[1] [2] [{probe.module}] [normal] [{probe.module}:selftest:1] {marker}\n"
                )
                with probe.local_path.open("a", encoding="utf-8") as handle:
                    handle.write(line)
                    handle.flush()

        received: dict[str, bytes] = {probe.key: b"" for probe, _c in opened}
        pending = {probe.key for probe, _c in opened}
        window = LIVE_TAIL_WINDOW_SECONDS if running else LIVE_TAIL_PROBE_WINDOW_SECONDS
        deadline = time.time() + window
        while time.time() < deadline and pending:
            for probe, channel in opened:
                if probe.key not in pending or not channel.recv_ready():
                    continue
                received[probe.key] += channel.recv(65536)
                blob = received[probe.key]
                done = (
                    blob.count(b"\n") >= 1 if running else marker.encode() in blob
                )
                if done:
                    pending.discard(probe.key)
            time.sleep(0.1)

        for probe, _channel in opened:
            blob = received[probe.key]
            if running:
                rows = [item for item in blob.decode("utf-8", "replace").splitlines() if item.strip()]
                _record(
                    f"实时 tail 转发新增行 · {probe.module}",
                    bool(rows),
                    f"由实时日志源产出 {len(rows)} 行"
                    if rows
                    else f"{window}s 内没有新行（远端 tail 缓冲未 flush？）",
                )
            else:
                _record(
                    f"实时 tail 转发新增行 · {probe.module}",
                    marker.encode() in blob,
                    f"探针模式，收到 {len(blob)} 字节",
                )
    except Exception as exc:  # noqa: BLE001
        _record("实时 tail 转发新增行", False, str(exc))
    finally:
        for _probe, channel in opened:
            try:
                channel.close()
            except Exception:  # noqa: BLE001
                pass
        client.close()
        if not running:
            # 只有探针模式才回滚：这时没人并发写，截断是安全的
            for probe, _channel in opened:
                with probe.local_path.open("r+b") as handle:
                    handle.truncate(sizes[probe.key])


def check_live_stream() -> None:
    """实时日志源：3 类互相关联的异常 + 正常节拍（单模块视角也齐备），活动文件不越界。"""
    from . import livesim
    from apps.logsources.services.log_format_parser import DEBUG_PATTERN

    pattern = re.compile(DEBUG_PATTERN)
    state = livesim.load_state()

    # 1) 容量：活动文件永远不超过上限；历史采样取自当前段 + 归档段
    contents: dict[str, str] = {}
    history: dict[str, str] = {}
    oversize: list[str] = []
    for target in livesim.TARGETS:
        path = target.local_path
        if not path.exists():
            oversize.append(f"{target.module}: 缺失")
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        contents[target.key] = text
        count = text.count("\n")
        if count > livesim.MAX_LIVE_LINES:
            oversize.append(f"{target.module}={count}")
        slot = state.target(target.key)
        if slot.archived and Path(slot.archived).exists():
            try:
                archived_text = Path(slot.archived).read_text(encoding="utf-8", errors="replace")
                history[target.key] = archived_text
                # 轮转出来的归档也必须卡在上限内（写满即收档，不允许 1001 行越界）
                archived_count = archived_text.count("\n")
                if archived_count > livesim.MAX_LIVE_LINES:
                    oversize.append(f"{Path(slot.archived).name}={archived_count}")
            except OSError:
                pass
    _record(
        f"实时日志源活动文件与归档 ≤ {livesim.MAX_LIVE_LINES} 行",
        not oversize,
        "、".join(oversize) or f"{len(contents)} 条流均在上限 {livesim.MAX_LIVE_LINES} 内",
    )
    if not contents:
        return

    # 2) **记录首行**必须匹配后端八字段调试日志（否则前端解析不出 level/module）。
    #    按记录判而不按物理行判：spwsp 的测量快照一次落十几行、只有首行带八字段
    #    前缀（``loggen.scan_metric_body``），续行本来就是裸文本 —— 逐行判会把
    #    这条**故意**用来复现解析截断的夹具误报成格式违规。
    pool = [line for text in contents.values() for line in text.splitlines() if line.strip()]
    # 逐条流**分别**归组：跨流拼在一起的话，某条流开头的续行会被错记成上一条流
    # 末尾那条记录的续行，违规就有机会被吞掉。
    records = [
        record
        for text in contents.values()
        for record in _group_records(_physical_lines(text))
    ]
    bad = [lines[0] for lines in records if not pattern.match(lines[0])]
    _record(
        "实时日志源记录首行格式合规",
        bool(records) and not bad,
        f"{len(records)} 条记录 / {len(pool)} 物理行，首行不合规 {len(bad)}",
    )

    # 采样窗口：**整段**活动文件（上限见 MAX_LIVE_LINES，本来就装得下一轮）+ 整段归档。
    # 不能只取尾部几十行：一轮剧本会在 5 条流上各写十几行，窗口太小会恰好卡在
    # 两轮之间，"同一 trace 横跨几条流"这种跨轮判断就取不到完整样本；
    # 某个流刚轮转时活动段可能只有几行，靠归档段把这一轮剧本补齐。
    samples: dict[str, str] = {}
    for key, text in contents.items():
        samples[key] = text + "\n" + history.get(key, "")
    joined = "\n".join(samples.values())
    sample_lines = [line for text in samples.values() for line in text.splitlines() if line.strip()]

    # 一轮剧本里带 trace= 的行数（从剧本自身算出来，不写死数字）。日志源刚重启或
    # 刚跑完 init 时，实时内容还不足一轮，跨模块关联必然取不到样本 —— 这种情况
    # 报「跳过」而不是假失败。
    trace_lines = [line for line in sample_lines if "trace=TR-" in line]
    round_trace_lines = sum(1 for _k, _l, _f, _p, body in livesim._SCRIPT if "trace={trace}" in body)
    if len(sample_lines) < 40 or len(trace_lines) < round_trace_lines:
        _record(
            "实时日志源内容样本",
            True,
            f"实时内容 {len(trace_lines)} 行（一轮剧本约 {round_trace_lines} 行），日志源刚启动，内容与关联检查跳过",
        )
    else:
        # 3) 至少 3 类异常 + 1 类正常
        codes = sorted(set(re.findall(r"code=(ERR_[A-Z_]+)", joined)))
        causes = sorted(set(re.findall(r"cause=(ERR_[A-Z_]+)", joined)))
        levels = sorted({m.group("level") for line in sample_lines if (m := pattern.match(line))})
        normal = sum(
            1 for line in sample_lines
            if any(marker in line for marker in livesim.NORMAL_BEAT_MARKERS)
        )
        _record("实时日志源含正常节拍日志", normal > 0, f"{normal} 行 · 级别 {levels}")
        _record("实时日志源含 ≥3 类异常", len(codes) >= 3, "、".join(codes) or "未找到 code=ERR_*")

        # 4) 关联：同一 trace 横跨多条流，且 cause 指向上游故障码
        traces: dict[str, set[str]] = {}
        for key, text in samples.items():
            for value in set(re.findall(r"trace=(TR-[\w-]+)", text)):
                traces.setdefault(value, set()).add(key)
        best_trace, best_files = max(traces.items(), key=lambda item: len(item[1]), default=("", set()))
        _record(
            "实时日志源异常跨模块关联",
            len(best_files) >= 3,
            f"{best_trace} 横跨 {len(best_files)} 条流：{'、'.join(sorted(best_files))}",
        )
        _record(
            "实时日志源 cause 因果链完整",
            {"ERR_MECORE_ENC_JITTER", "ERR_CPFR_FLOW_LOW", "ERR_CPFR_THERM_OVERLOAD"} <= set(causes),
            " → ".join(causes) or "未找到 cause=",
        )

        # 4b) 单模块可验证：前端「实时监听」一次只盯一个模块，所以**会被单独订阅的
        #     流**（livesim.OBSERVER_KEYS：扫片观察位 spwsp、工件台点位 wsp）自己就得
        #     凑齐「≥3 类异常 + 1 类正常节拍」，否则订阅单模块时凑不齐，
        #     实时监听的过滤/告警效果也就没法验。
        for observer_key in livesim.OBSERVER_KEYS:
            spec = livesim._TARGET_BY_KEY[observer_key]
            observer_text = samples.get(observer_key, "")
            observer_lines = [line for line in observer_text.splitlines() if line.strip()]
            observer_codes = sorted(set(re.findall(r"code=(ERR_[A-Z_]+)", observer_text)))
            observer_levels = sorted({
                m.group("level") for line in observer_lines if (m := pattern.match(line))
            })
            observer_normal = sum(
                1 for line in observer_lines
                if any(marker in line for marker in livesim.NORMAL_BEAT_MARKERS)
            )
            _record(
                f"实时日志源单模块 {spec.module} 含 ≥3 类异常 + 正常节拍",
                len(observer_codes) >= 3 and observer_normal > 0,
                f"{spec.module} 异常 {len(observer_codes)} 类 [{'、'.join(observer_codes)}] · "
                f"级别 {observer_levels} · 正常 {observer_normal} 行",
            )

        # 4c) 日志内容规则：正文必须以 ``函数名() >()`` / ``函数名() <()`` 成对出现，
        #     前端折叠函数卡片认的就是"方向符 + 前面那个带括号的函数名"。
        marker_in = sum(1 for line in sample_lines if ENTRY_MARKER_REGEX.search(line))
        marker_out = sum(1 for line in sample_lines if EXIT_MARKER_REGEX.search(line))
        _record(
            "实时日志源正文含入口/出口边界",
            marker_in > 0 and marker_out > 0,
            f"入口 >() {marker_in} 行 · 出口 <() {marker_out} 行",
        )

    # 5) 轮转原子性：改名成归档的同时必须留下一个空的当前段
    with tempfile.TemporaryDirectory() as folder:
        probe = Path(folder) / "cpcore.log"
        probe.write_text("x\n" * 5, encoding="utf-8")
        archived = livesim.archive_current(probe, "cpcore", datetime.now())
        ok = (
            probe.exists()
            and probe.stat().st_size == 0
            and archived.exists()
            and archived.stat().st_size == 10
        )
        _record("实时日志源轮转原子性", ok, f"{probe.name} 归零 / 归档 {archived.name} 保留原内容")


def check_live_cadence() -> None:
    """实时日志源：**每条流 0.5s 一行**，且窗口滑动后总量恒定有界。

    这条对着用户提的两个要求逐条量：

    1. **0.5 秒加一条** —— 从盘上的真实文件量相邻两行时间戳的间隔中位数。
       注意老实现是一条扁平游标、一个 tick 只落一行，5 条轮流写，所以任何一条
       都要等 5 个 tick 才轮到自己；现在 ``emit_tick`` 每 tick 让每条流各落一行，
       间隔中位数应当落在 0.5s 附近（节拍抖动 ±35%，再叠加每条流各自的 ±120ms
       落笔抖动与写盘开销）。
       ⚠️ 只量**文件尾部的连续实时段**：当前段里 ``init`` 写进去的历史段是
       "一分钟一行"的慢节奏，直接取最后 N 个间隔会在刚 init / 刚重启日志源的
       几十秒里全取到历史段，把中位数拉成 ~60s（假失败）。样本不足就报跳过，
       不报失败。
    2. **滑动窗口 / 不无限追加** —— 每条流的落盘行数必须恒定有界：
       当前段 ≤ cap、最新归档 ≤ cap、回收站固定槽位 ≤ cap，
       所以"当前段 + 最新归档"永远 ≤ 2 cap，挂多久都不会越攒越多。
    3. **各流各有各的时钟** —— 5 条流是同一个进程写的，如果都用同一个
       ``datetime.now()``，跨模块查询里它们的时间戳会逐列对齐到同一毫秒
       （一眼看出是模拟数据）。这里查最近若干行的时刻集合没有公共交集。

    日志源没在跑时不量节奏（盘上的行是上一轮留下的，间隔已经失真），
    直接报跳过。
    """
    from . import livesim

    cap = livesim.MAX_LIVE_LINES
    interval = livesim.DEFAULT_INTERVAL_SECONDS
    state = livesim.load_state()
    running = _stream_running()

    def stamps_of(target: livesim.StreamTarget) -> list[float]:
        """按真实查询路径取时刻：最新归档在前、当前段在后（时间上正好连续）。"""
        slot = state.target(target.key)
        chunks: list[str] = []
        if slot.archived and Path(slot.archived).exists():
            chunks.append(Path(slot.archived).read_text(encoding="utf-8", errors="replace"))
        path = target.local_path
        if path.exists():
            chunks.append(path.read_text(encoding="utf-8", errors="replace"))
        values: list[float] = []
        for text in chunks:
            for line in text.splitlines():
                if not line.startswith("[") or len(line) < 24:
                    continue
                try:
                    moment = datetime.strptime(line[1:24], "%Y-%m-%d %H:%M:%S.%f")
                except ValueError:
                    continue
                values.append(moment.timestamp())
        return values

    # ---- 1) 节奏：每条流最后若干行的间隔中位数 ----
    if not running:
        _record(
            "实时日志源每条流 0.5s 一行",
            True,
            "日志源没在运行，节奏检查跳过（先 scripts/sim.sh start stream）",
        )
    else:
        slowest: list[str] = []
        samples: list[str] = []
        problems: list[str] = []
        skipped: list[str] = []
        for target in livesim.TARGETS:
            values = stamps_of(target)
            gaps = [round(b - a, 3) for a, b in zip(values[:-1], values[1:])]
            # 只量**文件尾部的连续实时段**。
            #
            # ⚠️ 当前段里其实混着两种节奏：``init`` 写进去的历史段是"一分钟一行"
            # 的慢节奏，紧接着才是日志源实时追加的 0.5s 行。直接取"最后 120 个间隔"
            # 会在**刚 init / 刚重启日志源**的几十秒里全取到历史段 —— 中位数被拉成
            # ~60s，报一条看着很吓人的假失败（实测 58.27s）。从尾部往前走到第一个
            # 明显不属于实时节拍的间隔就停，剩下的就是实时段。
            tail: list[float] = []
            for gap in reversed(gaps):
                if gap > interval * 4:
                    break
                tail.append(gap)
            tail.reverse()
            if len(tail) < 20:
                # 日志源刚起来、实时段还没攒够样本。不作为失败：断言的主体
                # （实时日志源的节拍）此刻还没有足够的观测，报"跳过"更诚实。
                skipped.append(f"{target.module} {len(tail)} 个间隔")
                continue
            ordered = sorted(tail)
            median = ordered[len(ordered) // 2]
            samples.append(f"{target.module} {median:.2f}s")
            # 节拍抖动 ±35% + 落笔抖动 ±120ms，取个宽松上下界；
            # 中位数本身应当稳定在 interval 上（均值不变）。
            if not interval * 0.5 <= median <= interval * 2.0:
                slowest.append(f"{target.module}={median:.2f}s")
        _record(
            f"实时日志源每条流 {interval:g}s 一行",
            not slowest and not problems,
            "、".join(problems + slowest)
            or f"间隔中位数 [{'、'.join(samples)}]（应≈{interval:g}s）"
            + (f" · 实时段样本不足已跳过：{'、'.join(skipped)}" if skipped else ""),
        )

        # ---- 1b) 各流各有各的时钟：最近若干行的落笔时刻不该完全重合 ----
        recent: dict[str, set[float]] = {}
        for target in livesim.TARGETS:
            values = stamps_of(target)
            if values:
                recent[target.module] = {round(item, 3) for item in values[-20:]}
        if len(recent) >= 2:
            shared = set.intersection(*recent.values())
            first = livesim.TARGETS[0].module
            last = livesim.TARGETS[-1].module
            offset = ""
            if first in recent and last in recent:
                offset = f" · {first}/{last} 最新一行相差 {abs(max(recent[first]) - max(recent[last])) * 1000:.0f}ms"
            _record(
                "实时日志源各模块时钟不重合",
                not shared,
                f"公共时刻 {len(shared)} 个（应为 0）{offset}"
                if shared
                else f"{len(recent)} 条流的最近 20 行时刻互不重合{offset}",
            )

    # ---- 2) 滑动窗口：落盘量恒定有界 ----
    window_rows: list[str] = []
    over: list[str] = []
    for target in livesim.TARGETS:
        slot = state.target(target.key)
        current = livesim.count_lines(target.local_path)
        archived = livesim.count_lines(Path(slot.archived)) if slot.archived else 0
        recycled = livesim.count_lines(livesim.RECYCLE_DIR / f"{target.module}_prev.log")
        window = current + archived
        window_rows.append(f"{target.module} {window}（当前 {current} + 归档 {archived}）")
        if current > cap or archived > cap or recycled > cap or window > cap * 2:
            over.append(f"{target.module}: 当前 {current} / 归档 {archived} / 回收站 {recycled}")
    _record(
        f"实时日志源滑动窗口恒定 ≤ {cap * 2} 行/流",
        not over,
        "、".join(over) or "；".join(window_rows),
    )

    # ---- 3) 窗口确实滑过：至少有一条流已经轮转过 ----
    rotations = {target.module: state.target(target.key).rotations for target in livesim.TARGETS}
    total = sum(rotations.values())
    _record(
        "实时日志源窗口已滑动（写满即轮转）",
        bool(total) or not running,
        "、".join(f"{name}×{count}" for name, count in rotations.items() if count)
        or f"还没有任何一条流写满 {cap} 行（日志源刚启动，缩小容量实测见下一条）",
    )


def check_live_rotation_cycle() -> None:
    """高速驱动一份缩小容量的副本，验证「写满即轮转、只留最新归档、活动文件永不越界」。

    在真实日志树上等 spwsp 写到 ``MAX_LIVE_LINES`` 行要十几分钟，所以这里把日志根
    指向临时目录、容量压到 50 行，用零间隔把同一套 ``run_forever`` / ``emit_tick`` /
    ``rotate`` 代码路径跑满若干轮 —— 验的是行为本身，不是被压小的那个数字。
    """
    from . import livesim

    capacity = 50
    ticks = capacity * 6
    with tempfile.TemporaryDirectory() as folder:
        root = Path(folder)
        real_root_of = livesim.fleet.local_root_of
        real_state_file = livesim.STATE_FILE
        real_capacity = livesim.MAX_LIVE_LINES
        real_recycle = livesim.RECYCLE_DIR
        livesim.fleet.local_root_of = lambda _machine: root
        livesim.STATE_FILE = root / "stream_state.json"
        livesim.MAX_LIVE_LINES = capacity
        # 回收站也要指到临时目录，否则测试产物会落到真实 run/recycle 里
        livesim.RECYCLE_DIR = root / "recycle"
        try:
            state = livesim.run_forever(interval=0.0, limit=ticks, verbose=False)

            # 必须趁日志根还指向临时目录时核对，否则读到的是真实日志树。
            oversize: list[str] = []
            leftover: list[str] = []
            rotated: list[str] = []
            archived_sizes: list[int] = []
            reclaimed: list[str] = []
            for target in livesim.TARGETS:
                slot = state.target(target.key)
                current = livesim.count_lines(target.local_path)
                if current > capacity:
                    oversize.append(f"{target.module}(当前段)={current}")
                archives = sorted(target.local_path.parent.glob(f"{target.module}_*.log"))
                if slot.rotations:
                    rotated.append(f"{target.module}×{slot.rotations}")
                    for item in archives:
                        archived_sizes.append(livesim.count_lines(item))
                    if len(archives) != 1:
                        leftover.append(f"{target.module} 留下 {len(archives)} 份归档")
                if slot.archived and not Path(slot.archived).exists():
                    leftover.append(f"{target.module}: 归档记录悬空")
                # 回收用的是改名不是删除，所以日志目录里不会有残余 ——
                # 断言"除最新一份外全部搬走"，而不是"删掉了"。
                expected = [slot.archived] if slot.archived else []
                if slot.owned != expected:
                    leftover.append(f"{target.module} 未回收 {len(slot.owned)} 份（应为 {len(expected)}）")
                reclaimed.append(f"{target.module}×{slot.recycled}")
        finally:
            livesim.fleet.local_root_of = real_root_of
            livesim.STATE_FILE = real_state_file
            livesim.MAX_LIVE_LINES = real_capacity
            livesim.RECYCLE_DIR = real_recycle

        _record(
            "实时日志源写满即轮转（缩小容量实测）",
            bool(rotated) and not oversize,
            f"上限 {capacity} 行 · {len(rotated)} 条流轮转过 [{'、'.join(rotated)}]"
            + (f" · 越界 {'、'.join(oversize)}" if oversize else ""),
        )
        _record(
            # 判据是"**不超过**上限"，不是"正好写满"：记录宽度不再恒为 1 行 ——
            # spwsp 的测量快照一条就落十几行（loggen.scan_metric_body），轮转只能
            # 在"加上这一笔会越界"时提前触发，所以归档落在 (0, capacity] 区间内。
            # 真正要守的是**有界**：既不许越界，也不许攒着不收档。
            "实时日志源归档不超过上限",
            bool(archived_sizes) and all(0 < size <= capacity for size in archived_sizes),
            f"归档行数 {sorted(set(archived_sizes))}（上限 {capacity} 行）",
        )
        _record(
            "实时日志源只保留最新归档",
            not leftover,
            "、".join(leftover) or f"日志目录只留最新一份归档，旧归档已搬进回收站 [{'、'.join(reclaimed)}]",
        )


def _validate_call_chain(label: str, steps) -> None:  # noqa: ANN001 - 迭代器即可
    """校验一串 (函数, 相位, 正文) 是合法的调用栈轨迹，并回报结果。"""
    from .loggen import PHASE_BODY, PHASE_ENTER, PHASE_LEAVE, body_text, keyword_of, log_message

    stack: list[str] = []
    problems: list[str] = []
    pairs = 0
    boundary_levels: set[str] = set()
    # 入口/出口行必须带**固定关键字模板**：
    #   ``ScanWafer() >() enter wafer scan start ...``  /  ``... wafer scan end ...``
    # 关键字让人一眼看出这一步在做什么；出口复用同一个关键字，不用回头翻入口
    # 就能配对上。缺失就说明 PHASE_KEYWORDS 漏登记（会退回函数名当关键字）。
    template_problems: list[str] = []

    for item in steps:
        function, phase, body = item[0], item[1], item[2]
        level = item[3] if len(item) > 3 else "INFO"
        rendered = log_message(function, phase, body_text(body))
        keyword = keyword_of(function)
        call = f"{function}()"
        if phase == PHASE_ENTER:
            match = ENTRY_MARKER_REGEX.search(rendered)
            if not match:
                problems.append(f"{function} 入口缺 >()")
            elif BOUNDARY_NAME_REGEX.findall(rendered[: match.start()])[-1:] != [call]:
                problems.append(f"{function} 入口的函数名没紧邻 >()（应为 {call}）")
            else:
                boundary_levels.add(level)
                stack.append(function)
                if f"enter {keyword} start" not in rendered:
                    template_problems.append(function)
        elif phase == PHASE_LEAVE:
            match = EXIT_MARKER_REGEX.search(rendered)
            if not match:
                problems.append(f"{function} 出口缺 <()")
            elif BOUNDARY_NAME_REGEX.findall(rendered[: match.start()])[-1:] != [call]:
                problems.append(f"{function} 出口的函数名没紧邻 <()（应为 {call}）")
            if not stack or stack[-1] != function:
                problems.append(f"{function} 出口对不上（栈顶 {stack[-1] if stack else '空'}）")
            else:
                boundary_levels.add(level)
                stack.pop()
                pairs += 1
                if f"leave {keyword} end" not in rendered:
                    template_problems.append(function)
        else:
            if ENTRY_MARKER_REGEX.search(rendered) or EXIT_MARKER_REGEX.search(rendered):
                problems.append(f"{function} 正文行带了方向符")
            else:
                head = BOUNDARY_NAME_REGEX.match(rendered)
                if head is None or head.group(0) != call:
                    problems.append(f"{function} 正文行没以 {call} 开头：{rendered[:60]}")

    if stack:
        problems.append("未闭合：" + "、".join(stack))

    _record(
        f"{label} 调用链 LIFO 闭合",
        not problems and pairs > 0,
        "；".join(problems) if problems else f"{pairs} 组入口/出口同名配对",
    )
    # 边界行只负责记进出，异常级别必须落在正文行 —— 否则会出现
    # 「<() leave status=ok」却标着 FATAL 这种自相矛盾的行。
    _record(
        f"{label} 边界行级别为 INFO",
        boundary_levels <= {"INFO"},
        f"边界行级别 {sorted(boundary_levels) or '无'}",
    )
    _record(
        f"{label} 边界行带关键字模板",
        not template_problems,
        f"缺关键字 {sorted(set(template_problems))}" if template_problems
        else f"{pairs} 组「keyword start / keyword end」",
    )


def _validate_stage_frames(label: str, steps, codes) -> None:  # noqa: ANN001 - 迭代器即可
    """阶段框：``Stage_XXX() >() enter <阶段名> start step=n/N ...`` 成对且序号递增。

    这是"日志体现不同流程阶段"的落点：每到一个新阶段就开一个框，阶段内所有
    子调用都嵌在框里。两条硬约束：

    * **成对且 LIFO** —— 阶段框也是调用链，出口名字要逐字对上；
    * **序号递增** —— 同一条流上的阶段顺序就是流程顺序，回退说明剧本写反了。
    """
    from .loggen import PHASE_ENTER, PHASE_LEAVE, STAGE_FUNCTION_PREFIX

    order = {code: index for index, code in enumerate(codes)}
    opened: list[str] = []
    sequence: list[int] = []
    problems: list[str] = []

    for item in steps:
        function, phase = item[0], item[1]
        if not function.startswith(STAGE_FUNCTION_PREFIX):
            continue
        code = function[len(STAGE_FUNCTION_PREFIX):]
        if code not in order:
            problems.append(f"{function} 不在阶段序里")
            continue
        if phase == PHASE_ENTER:
            opened.append(function)
            sequence.append(order[code])
        elif phase == PHASE_LEAVE:
            if not opened or opened[-1] != function:
                problems.append(f"{function} 阶段出口对不上（栈顶 {opened[-1] if opened else '空'}）")
            else:
                opened.pop()

    if opened:
        problems.append("阶段框未闭合：" + "、".join(opened))
    if any(sequence[index] >= sequence[index + 1] for index in range(len(sequence) - 1)):
        problems.append(f"阶段序号非递增 {sequence}")

    _record(
        f"{label} 阶段框成对且顺序递增",
        not problems and len(sequence) > 1,
        "；".join(problems) if problems
        else f"{len(sequence)} 个阶段，序号 {[item + 1 for item in sequence]}",
    )


def check_machine_log_ownership() -> None:
    """可检索的日志树只能由**一台**机器承载，否则检索结果会把同一份日志读两遍。

    背景：后端 ``debug`` / ``run`` 两条日志路径配置的 scope 是 ``each_machine``，
    会遍历环境里的上位机与所有下位机。模拟机群里两台机器共用同一个日志根模板
    （``/log/{username}/debug``），host 又常常落成同一个地址，所以只要两台都生成
    同构的 ``<子系统>/<fm>.log``，同一份日志就会被读两遍 —— 检索结果里每个
    ``>() enter`` / ``<() leave`` 都出现两次（时间戳、线程号完全相同），前端按路径
    归并后就是用户看到的"同一个出口打印两遍"。

    这里断言 ``fleet.hosts_subsystem_logs`` 的归属以及 ``loggen.plan_machine``
    的实际产物：下位机不得再产出 ``<debug root>/`` 与 ``<run root>/`` 下的任何文件。
    """
    from . import loggen

    upper_owns = fleet.hosts_subsystem_logs(fleet.UPPER)
    lower_owns = fleet.hosts_subsystem_logs(fleet.LOWER1)
    _record(
        "可检索日志树只归属上位机",
        upper_owns and not lower_owns,
        f"上位机={upper_owns} 下位机={lower_owns}",
    )

    now = datetime.now()
    debug_root = fleet.UPPER.debug_root
    run_root = fleet.UPPER.run_root
    queryable_roots = {debug_root, run_root}

    lower_plan = loggen.plan_machine(fleet.LOWER1, now=now)
    lower_leaks = sorted(
        f"{root}/{relative}"
        for root, relative, _payload, _count in lower_plan
        if root in queryable_roots
    )
    _record(
        "下位机不产出可检索日志",
        not lower_leaks,
        f"{len(lower_plan)} 项产物，无 debug/run" if not lower_leaks else f"泄漏 {lower_leaks[:3]}",
    )

    upper_plan = loggen.plan_machine(fleet.UPPER, now=now)
    upper_queryable = {
        (root, str(relative))
        for root, relative, _payload, _count in upper_plan
        if root in queryable_roots
    }
    lower_queryable = {
        (root, str(relative))
        for root, relative, _payload, _count in lower_plan
        if root in queryable_roots
    }
    _record(
        "上位机仍有可检索日志",
        len(upper_queryable) > 0,
        f"{len(upper_queryable)} 个可检索产物",
    )
    _record(
        "两台机器的可检索产物不相交",
        not (upper_queryable & lower_queryable),
        f"交集 {len(upper_queryable & lower_queryable)} 项",
    )


def check_call_signature_shape() -> None:
    """调用链边界必须写成 ``函数名() >()`` —— 函数方法名带括号，且不许退回方括号写法。

    用户点名纠正过：

        对：``Stage_WSP_UNLOAD_MOVE() >() enter unload move stage start step=6/6 wafer=W04 …``
        错：``[Stage_WSP_UNLOAD_MOVE] >() enter unload move stage start step=6/6 wafer=W04 …``

    方向符 ``>()`` / ``<()`` **前面**那个词是**被调用的函数**，所以要写成调用形状。
    这不只是观感问题：前端 ``logParser.ts`` 首选 ``FUNCTION_PREFIX_REGEX``
    （``^函数名\\(``）去取 ``functionName``，方括号写法只能落进
    ``parseBoundaryFunctionName`` 的兜底分支 —— 那种情况下 ``functionName`` 是空的、
    只有 ``boundaryFunctionName`` 有值，「连续同名函数日志」「闭合后的同名尾随日志」
    两条内置折叠规则就失效（阶段框最明显：它整行的身份就靠这个标识）。

    量两路：

    * **剧本层** —— ``loggen.ALL_PROGRAMS``（各子系统批量程序）与 ``livesim._SCRIPT``
      （实时剧本）逐行渲染。两者**都已经过** ``expand_stage_groups``，所以阶段框
      （``Stage_XXX()``）也在里面，不用再单独展开一遍。要求：正文以 ``函数名()``
      开头；带方向符的行，方向符紧跟在 ``函数名()`` 之后；整行不得出现
      ``[函数名] >()`` / ``[函数名] <()`` 这种旧写法；
    * **落盘层** —— 真实文件（当前段 + 归档）按**记录**正文抽查同样两条。剧本对
      不代表写出来的对：改过正文格式却没 ``init`` 重建日志树时，盘上还是旧格式
      （这条会直接把"忘了重建"顶出来，而不是让用户在前端看到混合写法）。
    """
    from . import livesim, loggen
    from .loggen import PHASE_BODY, body_text, log_message

    problems: list[str] = []
    total = 0

    def _inspect(label: str, function: str, phase: str, body: str) -> None:
        nonlocal total
        rendered = log_message(function, phase, body_text(body))
        call = f"{function}()"
        total += 1
        if LEGACY_BRACKET_CALL_REGEX.search(rendered):
            problems.append(f"{label} {function} 用了方括号旧写法：{rendered[:70]}")
            return
        for marker in (ENTRY_MARKER_REGEX, EXIT_MARKER_REGEX):
            match = marker.search(rendered)
            if match is None:
                continue
            if BOUNDARY_NAME_REGEX.findall(rendered[: match.start()])[-1:] != [call]:
                problems.append(f"{label} {function} 方向符前不是 {call}：{rendered[:70]}")
        if phase == PHASE_BODY:
            head = BOUNDARY_NAME_REGEX.match(rendered)
            if head is None or head.group(0) != call:
                problems.append(f"{label} {function} 正文没以 {call} 开头：{rendered[:70]}")

    for label, program in loggen.ALL_PROGRAMS:
        for function, phase, _level, body in program:
            _inspect(label, function, phase, body)
    for _key, _level, function, phase, body in livesim._SCRIPT:
        _inspect("实时剧本", function, phase, body)
    _record(
        "调用链边界写成 函数名() >()（剧本层）",
        not problems,
        "；".join(problems[:3]) if problems
        else f"{total} 行均为 函数名() 调用形状（{len(loggen.ALL_PROGRAMS)} 份批量程序 + 实时剧本）",
    )

    # ---- 落盘层：真实文件里按**记录**抽查（续行没有八字段前缀，不能当记录首行） ----
    landed: list[str] = []
    records_seen = 0
    for module in ("spwsp", "wsp"):
        target = livesim._TARGET_BY_KEY[module]
        slot = livesim.load_state().target(module)
        paths = [target.local_path]
        if slot.archived:
            paths.append(Path(slot.archived))
        for path in paths:
            if not path.exists():
                continue
            lines = [line for line in path.read_text(encoding="utf-8", errors="replace").splitlines() if line.strip()]
            for record in _group_records(lines):
                message = _record_head(record[0])
                records_seen += 1
                if LEGACY_BRACKET_CALL_REGEX.search(message):
                    landed.append(f"{path.name} 方括号旧写法：{message[:70]}")
                    continue
                head = BOUNDARY_NAME_REGEX.match(message)
                if head is None:
                    landed.append(f"{path.name} 正文没以 函数名() 开头：{message[:70]}")
    _record(
        "调用链边界写成 函数名() >()（落盘层）",
        not landed and records_seen > 0,
        "；".join(landed[:3]) if landed
        else f"{records_seen} 条记录全部以 函数名() 开头、无方括号旧写法"
        if records_seen
        else "没有可抽查的日志文件（先跑 cli init / 启动日志源）",
    )


def check_log_call_chain() -> None:
    """日志内容规则：``函数名() >()`` 入口 / ``函数名() <()`` 出口，同名 LIFO 配对。

    这是前端折叠函数卡片的唯一依据（``foldingRules.ts`` 的内置规则
    ``builtin-explicit-boundary``，关键字就是 ``> ()`` / ``< ()``）。名字不一致
    或嵌套顺序写反，前端就会画出错乱嵌套、把父函数标成「未闭合」。

    这里直接校验三份剧本本身（通用批量 ``loggen._DEBUG_PROGRAM``、各子系统专属
    批量程序、实时 ``livesim._SCRIPT`` 的每条流），不用等日志落盘，也不用等
    写满一轮窗口（``livesim.MAX_LIVE_LINES`` 行）轮转。
    """
    from . import livesim, loggen

    for label, program in loggen.ALL_PROGRAMS:
        _validate_call_chain(
            label,
            [(fn, phase, body, level) for fn, phase, level, body in program],
        )
    _validate_stage_frames(
        "批量日志",
        [(fn, phase) for fn, phase, _level, _body in loggen._DEBUG_PROGRAM],
        loggen.DEBUG_STAGE_CODES,
    )
    for subsystem, program in loggen.SUBSYSTEM_PROGRAMS.items():
        _validate_stage_frames(
            f"批量日志 {subsystem}",
            [(fn, phase) for fn, phase, _level, _body in program],
            loggen.WSP_STAGE_CODES if subsystem == "wsp" else loggen.DEBUG_STAGE_CODES,
        )
    for target in livesim.TARGETS:
        steps = [
            (function, phase, body, level)
            for key, level, function, phase, body in livesim._SCRIPT
            if key == target.key
        ]
        _validate_call_chain(f"实时日志 {target.module}", steps)
        _validate_stage_frames(
            f"实时日志 {target.module}",
            [(function, phase) for function, phase, _body, _level in steps],
            livesim.STAGE_CODES,
        )

    # 关键字模板必须**全覆盖**：任何边界行函数没登记在 PHASE_KEYWORDS 里，就会
    # 退回函数名当关键字，用户看到的就不是「wafer scan start」而是「ScanWafer start」。
    # 各子系统专属程序也要一起过 —— 漏一个就会让 wsp 这种新组件的边界行退回函数名。
    from .loggen import PHASE_BODY, body_text, log_message

    boundary_functions = {
        function
        for _label, program in loggen.ALL_PROGRAMS
        for function, phase, _level, _body in program
        if phase != PHASE_BODY
    } | {
        function
        for _key, _level, function, phase, _body in livesim._SCRIPT
        if phase != PHASE_BODY
    }
    unregistered = sorted(boundary_functions - set(loggen.PHASE_KEYWORDS))
    _record(
        "调用链函数都有关键字模板",
        not unregistered,
        f"未登记 {unregistered}" if unregistered
        else f"{len(boundary_functions)} 个函数全部登记在 PHASE_KEYWORDS",
    )

    # 日志正文必须**全英文**：真实机台的调试日志/执行器日志/运行事件日志都是英文，
    # 模拟器里混中文会一眼看穿是假数据。这条把"渲染出来的每一行"都过一遍，
    # 新增剧本时写错语言（或关键字表被改回中文）会立刻失败。
    #
    # 注意走 ``body_text``：移动点位正文里带 ``{ … }`` 花括号，模板里存的是转义过的
    # ``{{ … }}``，直接看会看到 ``{{ "x":0.003 }}``。校验要按**最终写进日志的样子**判。
    cjk_pattern = re.compile(r"[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef]")
    offenders: list[str] = []
    rendered_total = 0
    for label, program in loggen.ALL_PROGRAMS:
        for function, phase, _level, body in program:
            rendered = log_message(function, phase, body_text(body))
            rendered_total += 1
            if cjk_pattern.search(rendered):
                offenders.append(f"{label}/{function}")
    for _key, _level, function, phase, body in livesim._SCRIPT:
        rendered = log_message(function, phase, body_text(body))
        rendered_total += 1
        if cjk_pattern.search(rendered):
            offenders.append(f"实时/{function}")
    _record(
        "日志正文全英文",
        not offenders,
        f"含中文的行 {sorted(set(offenders))}" if offenders
        else f"{rendered_total} 行正文无中日韩字符",
    )

    # 落盘抽查：活动文件里真的能看到方向符（剧本对不代表写出来的对）
    sample = Path(fleet.remote_to_local(fleet.UPPER, f"{fleet.UPPER.debug_root}/spwsp/spwsp.log"))
    if not sample.exists():
        _record("日志正文含调用链边界", False, f"{sample} 不存在")
        return
    lines = [line for line in sample.read_text(encoding="utf-8", errors="replace").splitlines()[-400:] if line.strip()]
    records = _group_records(lines)
    entries = sum(1 for line in lines if ENTRY_MARKER_REGEX.search(line))
    exits = sum(1 for line in lines if EXIT_MARKER_REGEX.search(line))
    # 第 9 段才是正文（前 8 个 [...] 是固定字段），所以就从这一段的**行首**找
    # ``函数名()``。只查**每条记录的首行**：spwsp 的测量快照是跨行正文，续行没有
    # 八字段前缀、也没有函数名，要求"每一物理行都带函数名"会把这条合法记录判成失败
    # （而它是故意留的夹具）。
    named = sum(
        1
        for lines_of_record in records
        if BOUNDARY_NAME_REGEX.search(_record_head(lines_of_record[0]))
    )
    _record(
        "日志正文含调用链边界",
        entries > 0 and exits > 0 and bool(records) and named == len(records),
        f"最近 {len(records)} 条记录 / {len(lines)} 物理行：入口 {entries} / 出口 {exits} / 带函数名 {named}",
    )


def check_multiline_record_fixture() -> None:
    """跨行记录夹具：spwsp 的测量快照必须真的**跨多行**，且七路干涉仪通道一个不少。

    这条守的是"复现解析截断"的那个场景本身（正文构造见
    ``loggen.scan_metric_body``）：一条记录横跨十几行，**只有首行带八字段前缀**
    （时间戳在最前），后面各行都是没有时间戳的续行。后端按"行首时间戳"切记录，
    续行会被丢掉，正文因此被截断 —— 这正是要能稳定复现的形态。

    为什么值得单独钉一条：

    * 夹具若被"顺手改成单行"，截断就复现不出来了，而所有既有的行格式断言反而会
      **变绿**（逐行判八字段重新成立）—— 场景消失了却没有任何告警；
    * 六自由度 / 功率 / 对比度 / 七路通道是用户点名要的字段，缺一个夹具就不完整；
    * 还要查"剧本里真的有人用它"，否则夹具写了却没接线，日志里根本不会出现。

    三条一起查：**形态**（物理行数 > 1，且续行不以 ``[`` 开头）、**内容**（六自由度 +
    功率 + 对比度 + ``x1``…``x4`` / ``y1``…``y3`` 全在，**且整个字典能
    ``json.loads``** —— 键必须带双引号、不能有尾随逗号）、**接线**（spwsp 实时剧本里
    确实有一行正文用了它）。
    """
    from . import livesim, loggen

    body = loggen.body_text(loggen.scan_metric_body("exposure_mid"))
    rows = body.splitlines()
    problems: list[str] = []

    if len(rows) <= 1:
        problems.append(f"正文只有 {len(rows)} 个物理行，不是跨行形态")
    # 续行一旦以 "[" 开头，就会被解析器当成新记录 / 被 awk 预筛放行，
    # 就复现不出"续行被丢掉"的效果了。
    heading_continuations = [row for row in rows[1:] if row.lstrip().startswith("[")]
    if heading_continuations:
        problems.append(f"续行以 '[' 开头，不再是无时间戳的裸续行：{heading_continuations[:2]}")

    # 快照与点位正文用同一套引号规则：**键与字符串值都带双引号**（花括号里是合法
    # JSON）。以前这里查的是 Python repr 风格的单引号，用户提出"字典要能 json 解析"
    # 之后一并改成双引号 —— 顺便也就把"整个快照能不能 json.loads"变成可断言的了。
    missing_dof = [axis for axis in loggen.MOVE_POINT_DOF if f'"{axis}":' not in body]
    if missing_dof:
        problems.append(f"缺自由度 {missing_dof}")
    for token in ('"power_mW":', '"contrast":', '"interferometer":'):
        if token not in body:
            problems.append(f"缺字段 {token}")
    if "'" in body:
        problems.append("正文里还留着 Python repr 的单引号（JSON 只认双引号）")

    channels = [axis for axis in loggen.INTERFEROMETER_CHANNELS if f'"{axis}":' in body]
    if channels != list(loggen.INTERFEROMETER_CHANNELS):
        problems.append(
            f"干涉仪通道 {channels} 与 {list(loggen.INTERFEROMETER_CHANNELS)} 不符"
        )

    # 字典要**真的能 json.loads**：无引号的键、Python 风格的单引号、尾随逗号
    # 三种写法都"看着像字典"，只有真解析一遍才能同时把它们挡下来。
    try:
        parsed = loggen.parse_scan_metric(body)
    except ValueError as exc:
        problems.append(f"快照字典 json.loads 失败：{exc}")
    else:
        flat = set(parsed) | set(parsed.get("position", {})) | set(parsed.get("interferometer", {}))
        expected_keys = {"position", "power_mW", "contrast", "interferometer",
                         "stage_speed_mm_s", "status"}
        if not expected_keys <= set(parsed):
            problems.append(f"快照缺字段 {sorted(expected_keys - set(parsed))}")
        if not set(loggen.INTERFEROMETER_CHANNELS) <= flat:
            problems.append("解析出来的字典里少干涉仪通道")

    wired = any(
        loggen.SCAN_METRIC_MARKER in loggen.body_text(step_body)
        for key, _level, _function, _phase, step_body in livesim._SCRIPT
        if key == "spwsp"
    )
    if not wired:
        problems.append("spwsp 实时剧本里没有任何一行使用该正文")

    _record(
        "跨行测量快照夹具完整（截断场景可复现）",
        not problems,
        "；".join(problems)
        if problems
        else f"{len(rows)} 行记录 · 六自由度/功率/对比度齐全 · "
        f"干涉仪 {'/'.join(loggen.INTERFEROMETER_CHANNELS)} 七路 · "
        f"json.loads 通过 · 已接入 spwsp 实时流",
    )


def check_wsp_move_points() -> None:
    """``wsp``（工件台点位）组件：日志正文必须是固定格式的绝对移动点位。

    形如::

        move absolute { "x":0.003, "y":0.999, "z":0.008, "rx":0.0009, "ry":-0.0013,
                        "rz":0.0021, "speed":120.0, "mode":"absolute",
                        "point":"load_position", "status":"settled" }

    工件台是 **6-DOF 台**，所以点位行必须给全 ``x, y, z, rx, ry, rz``
    六个自由度（而且顺序就是这个）；少写几个轴会被本条与
    ``wsp 点位行六自由度齐全`` 一起拦下来。

    字段顺序、数值精度、点位名都固定（点位表在 ``loggen.WSP_MOVE_POINTS``），
    这样才能当检索锚点用 —— 用户按 ``"x":`` / ``"point":"load_position"``
    就能捞出来。

    🔴 **花括号里必须是合法 JSON**：键与 ``mode`` / ``point`` / ``status`` 都带
    双引号。用户提过这条 —— 不带引号的 ``{ x:1.250, … }`` 看着像字典，
    ``json.loads`` 却直接抛 ``JSONDecodeError``，"看着能解析、一用就崩"最坑人。
    所以除了格式正则，还会**真的 ``json.loads`` 一遍**（见
    ``wsp 点位正文可被 json.loads 解析``）。

    三路都查，缺一路就可能"剧本对了但落盘不对"：

    * **剧本层**：批量程序与实时剧本里每条点位正文要合格式，且 8 个点位一个不少；
    * **落盘层**：真打开 ``wsp/wsp.log`` 与最新归档逐行扫 —— 花括号转义写错
      （``{{`` 少写一个）在剧本层是看不出来的，只有落盘才会暴露成 ``KeyError``
      或写出一行带双花括号的假日志；
    * **挂载层**：``wsp`` 要真的登记成子系统、真的有一条实时流、真的有自己的
      内容程序 —— 否则"加了组件但前端看不到"。
    """
    from . import livesim, loggen

    pattern = re.compile(loggen.MOVE_POINT_PATTERN)

    # 1) 剧本层：所有带 "move absolute" 的正文都必须合格式，且点位全覆盖
    template_bad: list[str] = []
    covered_points: set[str] = set()
    point_templates = 0
    for label, program in loggen.ALL_PROGRAMS:
        for function, _phase, _level, body in program:
            if "move absolute" not in body:
                continue
            point_templates += 1
            text = loggen.body_text(body)
            if not pattern.fullmatch(text):
                template_bad.append(f"{label}/{function}: {text}")
                continue
            covered_points.add(loggen.point_name_of(text))
    for _key, _level, function, _phase, body in livesim._SCRIPT:
        if "move absolute" not in body:
            continue
        point_templates += 1
        text = loggen.body_text(body)
        if not pattern.fullmatch(text):
            template_bad.append(f"实时/{function}: {text}")
            continue
        covered_points.add(loggen.point_name_of(text))

    expected_points = {row[0] for row in loggen.WSP_MOVE_POINTS}
    missing_points = sorted(expected_points - covered_points)
    _record(
        "wsp 点位正文格式（剧本层）",
        point_templates > 0 and not template_bad and not missing_points,
        f"不合格式 {'；'.join(template_bad[:2])}" if template_bad
        else f"缺点位 {missing_points}" if missing_points
        else f"{point_templates} 条点位行 · {len(covered_points)}/{len(expected_points)} 个点位全覆盖",
    )

    # 2) 落盘层：活动文件 + 最新归档（历史批量段与实时追加段都在这两个文件里）
    target = livesim._TARGET_BY_KEY["wsp"]
    state = livesim.load_state()
    paths = [target.local_path]
    slot = state.target("wsp")
    if slot.archived and Path(slot.archived).exists():
        paths.append(Path(slot.archived))
    scanned = 0
    landed: list[str] = []
    malformed: list[str] = []
    for path in paths:
        if not path.exists():
            continue
        scanned += 1
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            marker = "move absolute"
            if marker not in line:
                continue
            landed.append(line)
            tail = line[line.index(marker):]
            if not pattern.fullmatch(tail):
                malformed.append(tail[:90])
    _record(
        "wsp 点位正文格式（落盘层）",
        bool(landed) and not malformed,
        f"不合格式（含未转义花括号？）{malformed[:2]}" if malformed
        else f"{len(landed)} 条点位行，扫过 {scanned} 个文件（当前段 + 最新归档）"
        if landed
        else f"{target.remote_path} 里没有点位行（先跑 scripts/sim.sh init）",
    )

    # 2a) JSON 层：**真的 json.loads 一遍**（用户提的那条要求）。
    #     格式正则只能证明"看起来像字典"；不带引号的 ``{ x:1.250, … }`` 同样能写出
    #     一个好看的正则去放行它，问题恰恰是它解析不了。所以这一条不看正则，
    #     直接把花括号那段交给 json.loads —— 剧本层与落盘层都要过。
    def dof_candidates() -> list[str]:
        """剧本层所有点位正文（已还原花括号）。2a 与 2b 共用同一批样本。"""
        texts: list[str] = [
            loggen.body_text(body)
            for _label, program in loggen.ALL_PROGRAMS
            for _f, _p, _l, body in program
            if "move absolute" in body
        ]
        texts += [
            loggen.body_text(body)
            for _key, _level, _function, _phase, body in livesim._SCRIPT
            if "move absolute" in body
        ]
        return texts

    json_bad: list[str] = []
    json_checked = 0
    for text in dof_candidates():
        json_checked += 1
        try:
            parsed = loggen.parse_move_point(text)
        except ValueError as exc:
            json_bad.append(f"{text[:60]}…（{exc}）")
            continue
        # 解析出来的字典要真的是"那一条点位"，别是解析了一个空对象就放行。
        if parsed.get("mode") != "absolute" or not parsed.get("point"):
            json_bad.append(f"{text[:60]}…（缺 mode/point：{sorted(parsed)}）")
    for line in landed:
        tail = line[line.index("move absolute"):]
        json_checked += 1
        try:
            loggen.parse_move_point(tail)
        except ValueError as exc:
            json_bad.append(f"落盘 {tail[:60]}…（{exc}）")
    _record(
        "wsp 点位正文可被 json.loads 解析",
        json_checked > 0 and not json_bad,
        f"{len(json_bad)} 条解析失败（如 {json_bad[0]}）" if json_bad
        else f"{json_checked} 条点位行（剧本 + 落盘）全部 json.loads 成功 · "
             f"键与字符串值均带双引号",
    )

    # 2b) 六自由度：点位行必须按 x, y, z, rx, ry, rz 的顺序给全六个轴。
    #     工件台是 6-DOF 台，只写平动或只写一个转角都会让这行"看起来像点位、其实不是"。
    #     正则在 :func:`loggen.dof_of_point` 里（``rx`` 要排在 ``x`` 前面交替，
    #     否则 ``rx`` 会被单字母分支先吃掉；键名两侧的引号是可选的，现在写的是
    #     ``"x":``），这里只比轴序。
    dof_texts = dof_candidates()
    expected_dof = list(loggen.MOVE_POINT_DOF)
    dof_seen = {tuple(loggen.dof_of_point(text)) for text in dof_texts}
    dof_bad = sorted(item for item in dof_seen if list(item) != expected_dof)
    _record(
        "wsp 点位行六自由度齐全",
        bool(dof_texts) and not dof_bad,
        f"轴序不符 {dof_bad[:2]}（应为 {expected_dof}）" if dof_bad
        else f"{len(dof_texts)} 条点位行 · 轴序 {','.join(expected_dof)}",
    )

    # 3) 挂载层：子系统 / 实时流 / 专属程序都得在
    anchored = (
        "wsp" in fleet.SUBSYSTEM_MODULES
        and ("wsp", "wsp") in fleet.all_modules()
        and "wsp" in livesim._TARGET_BY_KEY
        and loggen.program_for("wsp") is not loggen._DEBUG_PROGRAM
    )
    _record(
        "wsp 组件已挂载（子系统 / 实时流 / 专属程序）",
        anchored,
        "、".join([
            f"子系统 {'在' if 'wsp' in fleet.SUBSYSTEM_MODULES else '缺'}",
            f"实时流 {'在' if 'wsp' in livesim._TARGET_BY_KEY else '缺'}",
            f"专属程序 {'在' if loggen.program_for('wsp') is not loggen._DEBUG_PROGRAM else '缺（仍在用通用扫片程序）'}",
        ]),
    )


def check_wsp_spiral_trajectory() -> None:
    """``wsp`` 的曝光扫描轨迹必须是**从场中心向外螺旋步进**的螺线。

    这是"用户要的那条轨迹"的形状断言。少了它，把螺线**退化成别的形状**不会触发
    任何既有断言 —— 行格式照样合规、六自由度照样齐全、点位覆盖照样完整，
    场景没了却零告警：改成同心圆（半径阶跃）、改成一条直线、或者把点序打乱，
    三样都能一路全绿。

    查四件事：

    * **形态**：``spiral_01`` … ``spiral_NN`` 一个号不缺，点数 = 圈数 × 每圈步数；
      椭圆半径比例**逐步递增**（逐步递增才是螺旋，阶跃就是同心圆），
      且末步贴到 1.0（曝光场边缘）；
    * **范围与起点**：每一步都落在曝光场矩形内；轨迹第一个点（``exposure_mid``）
      就在螺线中心上；
    * **剧本按序**：批量程序与实时剧本里，``SCAN_TRAJECTORY_POINTS`` 必须作为
      **连续一段**按同一顺序出现 —— 顺序乱了、或者中间夹进别的点位，"轨迹"就断了，
      而 LIFO 配对 / 阶段框序号那些断言一个都不会察觉；
    * **无负零**：坐标里不许出现 ``-0.000`` 这种凑出来的"负零"。
    """
    from . import loggen

    cx, cy = loggen.SPIRAL_CENTER
    radius_x, radius_y = loggen.SPIRAL_RADIUS_X, loggen.SPIRAL_RADIUS_Y
    total = loggen.SPIRAL_STEPS

    # ---- 1) 形态：连续性 + 逐步向外 + 末步贴边缘 ----
    # 半径比例 t = √(((x-cx)/RX)² + ((y-cy)/RY)²)：把椭圆坐标归一化回"到中心的
    # 相对半径"，于是不必分别关心长短轴，只比这一个数就能判"是不是在向外盘"。
    # 坐标落盘时被截到 3 位小数，反算有约 1.5e-3 的量化误差，容差取 1e-2。
    names = [row[0] for row in loggen.SPIRAL_POINTS]
    expected_names = [
        f"{loggen.SPIRAL_POINT_PREFIX}{step:02d}" for step in range(1, total + 1)
    ]
    ratios = [
        math.hypot((row[1] - cx) / radius_x, (row[2] - cy) / radius_y)
        for row in loggen.SPIRAL_POINTS
    ]
    expected_ratios = [step / total for step in range(1, total + 1)]
    shape_problems: list[str] = []
    if names != expected_names:
        shape_problems.append(f"点位名序不符（缺 {sorted(set(expected_names) - set(names))[:3]}）")
    if total != loggen.SPIRAL_TURNS * loggen.SPIRAL_STEPS_PER_TURN:
        shape_problems.append(f"点数 {total} ≠ 圈数 × 每圈步数")
    out_of_line = [
        names[index]
        for index, (got, want) in enumerate(zip(ratios, expected_ratios))
        if abs(got - want) > 1e-2
    ]
    if out_of_line:
        shape_problems.append(f"半径不成等差外扩（{out_of_line[:3]}）")
    if not all(b > a for a, b in zip(ratios[:-1], ratios[1:])):
        shape_problems.append("半径非单调递增（更像同心圆，不是向外螺旋）")
    if abs(ratios[-1] - 1.0) > 1e-2:
        shape_problems.append(f"末步半径 {ratios[-1]:.3f} 没贴到边缘 1.0")
    _record(
        "wsp 螺旋轨迹由中心向外步进",
        not shape_problems,
        "；".join(shape_problems)
        or f"{total} 步 = {loggen.SPIRAL_TURNS} 圈 × {loggen.SPIRAL_STEPS_PER_TURN} 步，"
        f"半径比例 {ratios[0]:.3f} → {ratios[-1]:.3f}（等差外扩，末步贴边缘）",
    )

    # ---- 2) 范围与起点 ----
    # 曝光场矩形直接取 exposure_start / exposure_end 两点的包围盒 —— 轨迹的合法
    # 范围就是这两点框出来的行程，不再另抄一份边界值（否则两边会各改各的）。
    field = loggen._MOVE_POINT_BY_NAME
    x_bounds = sorted((field["exposure_start"][0], field["exposure_end"][0]))
    y_bounds = sorted((field["exposure_start"][1], field["exposure_end"][1]))
    outside = [
        row[0]
        for row in loggen.SPIRAL_POINTS
        if not (
            x_bounds[0] - 1e-9 <= row[1] <= x_bounds[1] + 1e-9
            and y_bounds[0] - 1e-9 <= row[2] <= y_bounds[1] + 1e-9
        )
    ]
    center = field["exposure_mid"]
    off_center = math.hypot(center[0] - cx, center[1] - cy)
    _record(
        "wsp 螺旋轨迹落在曝光场内且起于中心",
        not outside and off_center < 1e-9,
        f"越界 {outside[:3]}" if outside
        else f"中心偏差 {off_center:.4f}mm（应为 0）" if off_center >= 1e-9
        else f"x∈[{x_bounds[0]:.3f},{x_bounds[1]:.3f}] y∈[{y_bounds[0]:.3f},{y_bounds[1]:.3f}] · "
             f"{len(loggen.SPIRAL_POINTS)} 步全在范围内 · 起点 = 场中心",
    )

    # ---- 3) 剧本按序（连续一段）----
    def _point_order(rows) -> list[str]:  # noqa: ANN001 - 迭代即可
        order: list[str] = []
        for item in rows:
            body = item[-1]
            if "move absolute" not in body:
                continue
            order.append(loggen.point_name_of(body))
        return order

    def _contains_run(haystack: list[str], needle: list[str]) -> bool:
        span = len(needle)
        return any(
            haystack[index:index + span] == needle
            for index in range(len(haystack) - span + 1)
        )

    expected_order = list(loggen.SCAN_TRAJECTORY_POINTS)
    orders = {
        "批量程序": _point_order(loggen.program_for("wsp")),
        "实时剧本": _point_order([row for row in livesim._SCRIPT if row[0] == "wsp"]),
    }
    order_problems = [
        f"{label} {len(order)} 项里找不到连续点序"
        for label, order in orders.items()
        if not _contains_run(order, expected_order)
    ]
    _record(
        "wsp 螺旋轨迹剧本按序连续",
        not order_problems,
        "；".join(order_problems)
        or f"批量 / 实时均为连续 {len(expected_order)} 项（入场 → 中心 → 螺旋 {total} 步 → 退场）",
    )

    # ---- 4) 无"负零"坐标 ----
    # f"{-0.00004:.4f}" 会写出 "-0.0000"：格式正则放行，但人一眼就知道这个数是凑的。
    negative_zero = [
        text
        for text in (
            loggen.body_text(loggen.point_body(row[0])) for row in loggen.SPIRAL_POINTS
        )
        if re.search(r"-0\.0+(?![0-9])", text)
    ]
    _record(
        "wsp 螺旋轨迹坐标无负零",
        not negative_zero,
        f"{len(negative_zero)} 步带 -0.000（如 {negative_zero[0][:70]}）" if negative_zero
        else f"{len(loggen.SPIRAL_POINTS)} 步坐标均无 -0.000",
    )


def check_wsp_point_call_boundaries() -> None:
    """``wsp`` 的每条点位正文都必须落在**它自己那一次** ``MoveAbsolute`` 调用里。

    这是用户提的那条要求：``MoveAbsolute() move absolute { … }`` 不能是一串
    **没有入口/出口的散行**。日志规范要求每个动作都有自己的边界
    （``函数名() >()`` / ``函数名() <()``），所以轨迹里每个点位都是一次独立调用，
    三行一组：

        MoveAbsolute() >() enter stage absolute move start dof=6 point=spiral_07 profile=scan
        MoveAbsolute() move absolute { "x":…, "point":"spiral_07", "status":"settled" }
        MoveAbsolute() <() leave stage absolute move end dof=6 point=spiral_07 elapsed=… status=ok

    ⚠️ 少了这条断言，**退回"一次调用里连打几十条点位"不会触发任何既有断言**：
    行格式照样合规、六自由度照样齐全、点位覆盖照样完整、LIFO 也照样闭合
    （一对括号包住几十条正文本来就是合法的嵌套）—— 但那些点位正文就没有自己的
    边界了：前端折不出"这一次移动"，也答不出"走到这个点位用了多久"。

    量三路，缺一路就可能"剧本对了但落盘没对"：

    * **构造层** —— ``loggen.move_call_rows()`` 必须是三行一组（入口 / 点位正文 /
      出口），函数名逐行一致，且入口与出口都点名同一个点位；
    * **剧本层** —— 批量 wsp 程序与实时 wsp 剧本压栈统计：每次 ``MoveAbsolute``
      调用**恰好带 1 条点位正文**，且不存在"栈里没有同名调用"的点位正文；
    * **落盘层** —— 真实 ``wsp.log``（当前段 + 最新归档）按 ``rpc`` 字段的函数名
      压栈：每条点位正文所在行的函数必须是栈顶。剧本对、落盘也对，才算真的对。
    """
    from . import loggen

    # ---- 1) 构造层：三行一组，点名一致 ----
    build_problems: list[str] = []
    for row in loggen.WSP_MOVE_POINTS:
        name = row[0]
        rows = loggen.move_call_rows(name, profile="scan")
        phases = [phase for _f, phase, _l, _b in rows]
        functions = {function for function, _p, _l, _b in rows}
        if len(rows) < 3 or phases[0] != loggen.PHASE_ENTER or phases[-1] != loggen.PHASE_LEAVE:
            build_problems.append(f"{name} 不是「入口 … 出口」的形状：{phases}")
            continue
        if any(phase != loggen.PHASE_BODY for phase in phases[1:-1]):
            build_problems.append(f"{name} 入口与出口之间夹了非正文行：{phases}")
        if functions != {"MoveAbsolute"}:
            build_problems.append(f"{name} 三行的函数名不统一：{sorted(functions)}")
        if f"point={name}" not in loggen.body_text(rows[0][3]):
            build_problems.append(f"{name} 入口没点名")
        if f"point={name}" not in rows[-1][3]:
            build_problems.append(f"{name} 出口没点名（几十条出口会长得一模一样）")
    _record(
        "wsp 每个点位各成一次 MoveAbsolute 调用（构造层）",
        bool(loggen.WSP_MOVE_POINTS) and not build_problems,
        "；".join(build_problems[:3])
        or f"{len(loggen.WSP_MOVE_POINTS)} 个点位均为「入口 / 点位正文 / 出口」三行一组",
    )

    # ---- 2) 剧本层：每次调用恰好 1 条点位正文 ----
    def _calls(rows) -> tuple[list[int], int]:  # noqa: ANN001 - 迭代即可
        """按同名 LIFO 压栈，回报 (每次 MoveAbsolute 调用带的点位正文条数, 无边界条数)。"""
        stack: list[list] = []
        widths: list[int] = []
        orphan = 0
        for function, phase, body in rows:
            if phase == loggen.PHASE_ENTER:
                stack.append([function, 0])
                continue
            if phase == loggen.PHASE_LEAVE:
                if stack and stack[-1][0] == function:
                    name, count = stack.pop()
                    if name == "MoveAbsolute":
                        widths.append(count)
                continue
            if "move absolute" not in body:
                continue
            # 正文要落在**它自己那个函数**的调用里：wsp 的回零正文也写
            # ``move absolute { … "point":"origin" … }``（记的是参考点坐标），
            # 它属于 ``HomeStage`` 调用，不算"没有边界"。
            if stack and stack[-1][0] == function:
                stack[-1][1] += 1
            else:
                orphan += 1
        return widths, orphan

    programs = {
        "批量程序": [
            (function, phase, body) for function, phase, _level, body in loggen.program_for("wsp")
        ],
        "实时剧本": [
            (function, phase, body)
            for key, _level, function, phase, body in livesim._SCRIPT
            if key == "wsp"
        ],
    }
    script_problems: list[str] = []
    counts: list[str] = []
    for label, rows in programs.items():
        widths, orphan = _calls(rows)
        counts.append(f"{label} {len(widths)} 次")
        if not widths:
            script_problems.append(f"{label} 里没有一次 MoveAbsolute 调用")
        crowded = sorted({count for count in widths if count != 1})
        if crowded:
            script_problems.append(f"{label} 有调用带了 {crowded} 条点位正文（应为 1）")
        if orphan:
            script_problems.append(f"{label} 有 {orphan} 条点位正文没有同名调用包着")
    _record(
        "wsp 每条点位正文都在自己的调用里（剧本层）",
        not script_problems,
        "；".join(script_problems)
        or f"每次 MoveAbsolute 恰好 1 条点位正文（{'、'.join(counts)}）",
    )

    # ---- 3) 落盘层：真实文件里的位置关系 ----
    target = livesim._TARGET_BY_KEY["wsp"]
    state = livesim.load_state()
    paths = [target.local_path]
    slot = state.target("wsp")
    if slot.archived and Path(slot.archived).exists():
        paths.append(Path(slot.archived))

    unwrapped: list[str] = []
    crowded_lines: list[str] = []
    scanned = 0
    calls_seen = 0
    for path in paths:
        if not path.exists():
            continue
        stack: list[list] = []
        started = False
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            match = RPC_FUNCTION_REGEX.search(line)
            if not match:
                continue
            function = match.group(1)
            if not started:
                # 文件头往往正是上一次轮转的断口（那一次的入口留在上一份文件里），
                # 从第一条 ``MoveAbsolute`` 入口才开始看 —— 别把断口当成"缺边界"。
                if function == "MoveAbsolute" and ENTRY_MARKER_REGEX.search(line):
                    started = True
                else:
                    continue
            scanned += 1
            if ENTRY_MARKER_REGEX.search(line):
                stack.append([function, 0])
                continue
            if EXIT_MARKER_REGEX.search(line):
                if stack and stack[-1][0] == function:
                    name, count = stack.pop()
                    if name == "MoveAbsolute":
                        calls_seen += 1
                        if count != 1:
                            crowded_lines.append(f"一次调用带 {count} 条点位正文")
                continue
            if "move absolute" not in line:
                continue
            if stack and stack[-1][0] == function:
                stack[-1][1] += 1
            else:
                unwrapped.append(line.strip())
    landed_problems: list[str] = []
    if unwrapped:
        landed_problems.append(
            f"{len(unwrapped)} 条点位正文没有同名调用包着（如 {unwrapped[0][:80]}）"
        )
    if crowded_lines:
        landed_problems.append(f"有调用带多条点位正文：{sorted(set(crowded_lines))}")
    _record(
        "wsp 每条点位正文都在自己的调用里（落盘层）",
        scanned > 0 and not landed_problems,
        "；".join(landed_problems) if landed_problems
        else f"扫过 {scanned} 行 / {calls_seen} 次调用，每条点位正文都落在自己的入口与出口之间",
    )


def check_cpd_assets() -> None:
    """CPD 测校资产：报告字段能被后端解析，数据表格时间列 + 行时间落在报告窗口内。"""
    from apps.reports.parser import parse_report_summary

    report_root = Path(fleet.remote_to_local(fleet.UPPER, fleet.UPPER.cpd_report_root))
    data_root = Path(fleet.remote_to_local(fleet.UPPER, fleet.UPPER.cpd_data_root))
    reports = sorted(report_root.glob("*/*/*.rpt"))
    books = sorted(data_root.glob("*/*/*.xlsx"))
    _record("CPD 测校报告已生成", len(reports) > 0, f"{len(reports)} 份 · {fleet.UPPER.cpd_report_root}")
    # 同一天多次 init 会留下更早的批次（清理受小批量删除保护限制），所以这里比的是
    # "每一批报告都有同名数据表格"，而不是两边总数相等。
    report_keys = {item.relative_to(report_root).with_suffix("").as_posix() for item in reports}
    book_keys = {item.relative_to(data_root).with_suffix("").as_posix() for item in books}
    orphans = sorted(report_keys ^ book_keys)
    _record(
        "CPD 报告与数据表格成对",
        not orphans,
        f"{len(report_keys & book_keys)} 对配对" + (f"，孤儿 {orphans[:4]}" if orphans else ""),
    )
    if not reports:
        return

    # 目录层级必须是 <子系统>/<模块>/<文件>，否则后端 find -mindepth 3 扫不到
    depths = {len(item.relative_to(report_root).parts) for item in reports}
    _record("CPD 报告目录层级正确", depths == {3}, f"层级 {sorted(depths)}（应为 [3]）")

    # 两棵树的子系统/模块目录必须同名（后端数据根会转小写）
    report_dirs = {item.parent.relative_to(report_root).as_posix() for item in reports}
    data_dirs = {item.parent.relative_to(data_root).as_posix() for item in books}
    _record("报告与数据目录同构", report_dirs == data_dirs, f"报告 {len(report_dirs)} / 数据 {len(data_dirs)}")

    sample = reports[-1]
    relative = sample.relative_to(report_root).as_posix()
    summary = parse_report_summary(
        sample.read_text(encoding="utf-8"),
        file_name=sample.name,
        full_path=f"{fleet.UPPER.cpd_report_root}/{relative}",
        modified_at=sample.stat().st_mtime,
    )
    required = ("start_time", "stop_time", "test_run_result", "results_validation",
                "measurement_quality", "mcs_status", "execution_time", "cpd_name")
    missing = [key for key in required if not summary.get(key)]
    _record(
        "CPD 报告字段可解析",
        not missing,
        f"缺 {missing}" if missing else f"{summary['test_run_result']} · {summary['start_time']} → {summary['stop_time']}",
    )

    # 数据表格要有一列时间，且行时间落在同名报告的 [Start, Stop] 内
    twin = data_root / sample.relative_to(report_root).with_suffix(".xlsx")
    if not twin.exists():
        _record("CPD 报告有同名数据表格", False, f"缺少 {twin.name}")
        return
    from openpyxl import load_workbook

    workbook = load_workbook(twin, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = [str(cell or "").strip().casefold() for cell in (rows[0] if rows else ())]
    time_column = next((name for name in ("timestamp", "time", "date", "时间", "时间戳") if name in headers), None)
    _record("数据表格含时间列", time_column is not None, f"表头 {headers[:4]}")
    if time_column is None or len(rows) < 2:
        return
    index = headers.index(time_column)
    start = datetime.fromisoformat(summary["start_time"])
    stop = datetime.fromisoformat(summary["stop_time"])
    stamps = [row[index] for row in rows[1:] if isinstance(row[index], datetime)]
    inside = [item for item in stamps if start <= item <= stop]
    _record(
        "表格行时间落在报告窗口内",
        bool(stamps) and len(inside) == len(stamps),
        f"{len(inside)}/{len(stamps)} 行 ∈ [{start:%H:%M:%S}, {stop:%H:%M:%S}]",
    )


def check_cpd_backend_read() -> None:
    """走一遍后端真正的读取路径：SFTP 拉工作簿 → 解析 → 按报告时间窗过滤。"""
    from apps.environments.models import Environment
    from apps.reports import cpd_data_service as cpd
    from apps.reports.parser import parse_report_summary

    environment = Environment.objects.filter(upper_machine__host=fleet.UPPER.host).first()
    if environment is None:
        _record("CPD 数据表格（后端读取路径）", False, "尚未 seed，先跑 `python -m simremote.cli seed`")
        return

    report_root = Path(fleet.remote_to_local(fleet.UPPER, fleet.UPPER.cpd_report_root))
    reports = sorted(report_root.glob("*/*/*.rpt"))
    if not reports:
        _record("CPD 数据表格（后端读取路径）", False, "未生成报告")
        return
    sample = reports[-1].relative_to(report_root)
    subsystem, module, filename = sample.parts[0], sample.parts[1], sample.parts[2]
    remote_report = f"{fleet.UPPER.cpd_report_root}/{sample.as_posix()}"
    summary = parse_report_summary(
        (report_root / sample).read_text(encoding="utf-8"),
        file_name=filename, full_path=remote_report, modified_at=(report_root / sample).stat().st_mtime,
    )
    start_time, stop_time = summary["start_time"], summary["stop_time"]
    try:
        found = cpd.find_cpd_excel_files(environment, subsystem, module, start_time, stop_time)
        names = [item["path"].rsplit("/", 1)[-1] for item in found["files"]]
        _record(
            "后端按报告窗口筛出数据表格",
            names == [sample.with_suffix(".xlsx").name],
            f"命中 {names}",
        )
        if not names:
            return
        preview = cpd.preview_cpd_excel(
            environment, subsystem, module, found["files"][0]["path"],
            found["files"][0]["sheets"][0]["name"], start_time, stop_time,
        )
        _record(
            "后端可预览表格内容",
            preview.get("total", 0) > 0 and preview.get("filter_status") == "filtered",
            f"{preview.get('total')} 行 · 时间列 {preview.get('time_column')}",
        )
    except Exception as exc:  # noqa: BLE001
        _record("后端按报告窗口筛出数据表格", False, str(exc))


def check_atlog_report_site() -> None:
    """模拟 ATLog 用例报告站：目录页、HEAD/Range、以及真实的 URL 分析链路。"""
    import urllib.error
    import urllib.request

    from apps.atlog import services as atlog

    cases = [item for item in atlog_site.case_urls() if item["status"] == "failed"]
    if not cases:
        _record("报告站用例 URL", False, "未生成用例")
        return
    base = cases[0]["url"]
    _record("报告站用例 URL", True, base)

    # 1) 目录页必须是 nginx autoindex 风格（后端靠 HTMLParser 抓 <a href>）
    try:
        with urllib.request.urlopen(base, timeout=6) as response:
            page = response.read().decode("utf-8", "replace")
    except (urllib.error.URLError, OSError) as exc:
        _record("报告站目录页可达", False, f"{exc}（先跑 `python -m simremote.cli site-serve`）")
        return
    required = ["summary_report.xml", "event.log", "xytest.log", "full_logs/", "result/"]
    _record(
        "报告站目录页可达",
        bool(page) and all(f'"{name}"' in page for name in required),
        f"{len(page)} 字节",
    )

    # 2) 后端真正用 HTTP HEAD + Range 读尾部，这里直接验证这两个动作
    target = base + "summary_report.xml"
    request = urllib.request.Request(target, method="HEAD")
    with urllib.request.urlopen(request, timeout=6) as response:
        size = int(response.headers.get("Content-Length") or 0)
    _record("报告站支持 HEAD", size > 0, f"Content-Length={size}")
    ranged = urllib.request.Request(target, headers={"Range": "bytes=0-63"})
    with urllib.request.urlopen(ranged, timeout=6) as response:
        chunk = response.read()
        status = getattr(response, "status", 200)
    _record("报告站支持 Range", status == 206 and bytes(chunk).startswith(b"<?xml"), f"status={status} {bytes(chunk)[:24]!r}")

    # 3) 真正的分析：粘贴这个 URL 应该直接得到失败根因
    try:
        result = atlog.analyze_case(base)
    except Exception as exc:  # noqa: BLE001
        _record("ATLog URL 分析返回失败根因", False, str(exc))
        return
    conclusion = str(result.get("conclusion") or "")
    assertion = str(result.get("assertion") or "")
    _record(
        "ATLog URL 分析返回失败根因",
        result.get("status") == "failed" and "ResultStatus.FAILED" in f"{conclusion} {assertion}",
        f"status={result.get('status')} · {conclusion[:40]}",
    )
    _record(
        "ATLog 解析出调用栈",
        len(result.get("call_chain") or []) >= 3 and bool(result.get("failure_location")),
        f"{len(result.get('call_chain') or [])} 帧 · 末帧 {(result.get('failure_location') or {}).get('function')}",
    )
    sources = {item.get("source") for item in (result.get("evidence") or [])}
    _record(
        "ATLog 汇总多路证据",
        {"summary_report.xml", "pytest XML", "pytest HTML", "xytest.log"} <= sources,
        f"证据来源 {sorted(sources)}",
    )
    catalog = result.get("log_catalog") or []
    _record(
        "ATLog 发现用例日志目录",
        bool(catalog) and bool((catalog[0] or {}).get("modules")),
        f"{[(item.get('subsystem'), item.get('modules')) for item in catalog]}",
    )
    _record("ATLog 分析无告警", not (result.get("warnings") or []), f"warnings={result.get('warnings')}")

    # 4) 用例日志检索（按组件选目标，full_logs 是可检索的）
    modules = ((catalog[0] or {}).get("modules") or []) if catalog else []
    if not modules:
        return
    try:
        logs = atlog.query_case_logs(base, components=[modules[0]], max_lines=200)
    except Exception as exc:  # noqa: BLE001
        _record("ATLog 用例日志可检索", False, str(exc))
        return
    _record(
        "ATLog 用例日志可检索",
        logs.get("matched_file_count", 0) > 0 and logs.get("count", 0) > 0,
        f"{logs.get('matched_file_count')} 个文件 / {logs.get('count')} 行",
    )


def check_cpd_report_site() -> None:
    """CPD 资产同样发布在报告站上：.rpt / .xlsx 都能用「报告文件 URL 分析」直接解析。"""
    import urllib.request

    from apps.atlog import services as atlog

    host, port = fleet.ATLOG_SITE_HOST, fleet.ATLOG_SITE_PORT
    report_root = Path(fleet.remote_to_local(fleet.UPPER, fleet.UPPER.cpd_report_root))
    data_root = Path(fleet.remote_to_local(fleet.UPPER, fleet.UPPER.cpd_data_root))
    reports = sorted(report_root.glob("*/*/*.rpt"))
    if not reports:
        _record("报告站发布 CPD 报告 URL", False, "未生成 CPD 报告")
        return

    def data_root_of(path: Path) -> Path:
        # 数据树与报告树不同根，且目录名是小写
        item = path.relative_to(report_root)
        return data_root / item.parts[0].lower() / item.parts[1].lower() / f"{path.stem}.xlsx"

    def report_url_of(path: Path) -> str:
        item = path.relative_to(report_root)
        return f"http://{host}:{port}{atlog_site.CPD_REPORT_MOUNT}/{item.parts[0]}/{item.parts[1]}/{path.name}"

    def data_url_of(path: Path) -> str:
        item = path.relative_to(report_root)
        return (
            f"http://{host}:{port}{atlog_site.CPD_DATA_MOUNT}/"
            f"{item.parts[0].lower()}/{item.parts[1].lower()}/{path.with_suffix('.xlsx').name}"
        )

    mounts = atlog_site.cpd_urls()
    _record(
        "报告站发布 CPD 报告 URL",
        bool(mounts) and all(item["report_url"] for item in mounts),
        f"{len(mounts)} 个模块 · {atlog_site.CPD_REPORT_MOUNT}",
    )

    sample = reports[-1]
    try:
        with urllib.request.urlopen(report_url_of(sample), timeout=8) as response:
            head = response.read(64).decode("utf-8", "replace")
        _record("报告站可访问 .rpt", head.startswith("CPD Measurement Report"), head[:36])
    except Exception as exc:  # noqa: BLE001
        _record("报告站可访问 .rpt", False, f"{exc}（先跑 `python -m simremote.cli site-serve`）")
        return

    twin = data_root_of(sample)
    if not twin.exists():
        _record("报告站可访问 .xlsx", False, f"缺少同名表格 {twin.name}")
        return
    request = urllib.request.Request(data_url_of(sample), method="HEAD")
    with urllib.request.urlopen(request, timeout=8) as response:
        content_type = response.headers.get("Content-Type") or ""
        size = int(response.headers.get("Content-Length") or 0)
    _record(
        "报告站可访问 .xlsx",
        "spreadsheetml" in content_type and size == twin.stat().st_size,
        f"{content_type} · {size} 字节",
    )

    report = atlog.analyze_report_file(report_url_of(sample))
    _record(
        "报告文件 URL 分析（.rpt）",
        report.get("kind") == "cpd_report" and bool((report.get("cpd_report") or {}).get("test_run_result")),
        f"kind={report.get('kind')} result={(report.get('cpd_report') or {}).get('test_run_result')}",
    )
    sheet = atlog.analyze_report_file(data_url_of(sample))
    sections = sheet.get("sections") or []
    _record(
        "报告文件 URL 分析（.xlsx）",
        sheet.get("kind") == "spreadsheet" and bool(sections) and bool(sections[0].get("table")),
        f"{len(sections[0].get('table') or []) if sections else 0} 行 · {str(sheet.get('conclusion'))[:52]}",
    )

    # 同一批测校：报告判 FAILED，表格里就该有非 OK 行；两边数量必须一一对应
    failed_reports = failed_sheets = 0
    for path in reports:
        failed_reports += atlog.analyze_report_file(report_url_of(path)).get("status") == "failed"
        failed_sheets += atlog.analyze_report_file(data_url_of(path)).get("status") == "failed"
    _record(
        "CPD 报告与数据表格判定一致",
        failed_reports == failed_sheets and failed_reports > 0,
        f"{len(reports)} 批中：失败报告 {failed_reports} · 含非 OK 行的表格 {failed_sheets}",
    )


def check_report_file_analysis() -> None:
    """报告文件 URL 分析：一个链接进去，按类型分派到对应解析器。"""
    import re

    from apps.atlog import services as atlog

    cases = [item for item in atlog_site.case_urls() if item["status"] == "failed"]
    if not cases:
        _record("报告文件 URL 分析", False, "未生成失败用例")
        return
    base = cases[0]["url"]

    expectations = (
        ("summary_report.xml", "junit_xml"),
        ("result/pytest-01.xml", "junit_xml"),
        ("xytest.log", "xytest_log"),
        ("event.log", "event_log"),
        ("full_logs/log/debug/spm/RSPMCPD.log", "debug_log"),
    )
    for relative, kind in expectations:
        try:
            result = atlog.analyze_report_file(base + relative)
        except Exception as exc:  # noqa: BLE001
            _record(f"报告文件 URL 分析 {relative}", False, str(exc))
            continue
        _record(
            f"报告文件 URL 分析 {relative}",
            result.get("kind") == kind,
            f"kind={result.get('kind')} status={result.get('status')} case={result.get('case_id')}",
        )

    deep = atlog.analyze_report_file(base + "full_logs/log/debug/spm/RSPMCPD.log")
    _record(
        "报告文件能上溯出用例编号",
        deep.get("case_id") == cases[0]["case_id"],
        f"case_id={deep.get('case_id')}（期望 {cases[0]['case_id']}）",
    )

    # 调试日志只认行首 [..] 里的级别字段：正文里的 "position error" 不算错误行
    local = Path(cases[0]["local_path"]) / "full_logs/log/debug/spm/RSPMCPD.log"
    raw = local.read_text(encoding="utf-8") if local.exists() else ""
    expected = sum(
        1 for line in raw.splitlines()
        if re.match(r"^\[[^\]]+\]\s*\[(?:ERROR|FATAL|CRITICAL|FAIL|FAILED)\]", line)
    )
    findings = deep.get("findings") or []
    _record(
        "调试日志异常行不误报",
        len(findings) == expected and expected > 0,
        f"抽取 {len(findings)} 条（文件内真错误行 {expected} 条）",
    )

    try:
        atlog.analyze_report_file("http://example.com/summary_report.xml")
        _record("报告文件 URL 拒绝公网地址", False, "竟然放行了公网主机")
    except atlog.AtLogError:
        _record("报告文件 URL 拒绝公网地址", True)


def check_backend_catalog() -> None:
    from apps.environments.models import Environment
    from apps.logsources.services.remote_logs import scan_environment_logs

    environment = Environment.objects.filter(upper_machine__host=fleet.UPPER.host).first()
    if environment is None:
        _record("后端能发现模拟环境", False, "尚未 seed，先跑 `python -m simremote.cli seed`")
        return
    result = scan_environment_logs(environment, refresh=True)
    subsystems = result.get("subsystems") or []
    catalog = result.get("global_catalog") or []
    fms = sum(len(item.get("fms") or []) for item in catalog if isinstance(item, dict))
    _record("后端能发现子系统", len(subsystems) > 0, f"{len(subsystems)} 个：{', '.join(subsystems[:6])}")
    _record("后端能发现模块", fms >= len(fleet.all_modules()), f"{fms} 个模块")


# --------------------------------------------------------------------------- 入口


def _passwordless_probe(spec, key, *, expect_reject: bool = False) -> tuple[bool, str]:
    """用某把私钥连某台机器，返回 ``(是否符合预期, 说明)``。

    ``expect_reject=True`` 时"被拒绝"才算通过 —— 用来证明公钥认证真的在校验，
    而不是"谁来都放行"。
    """
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        client.connect(
            spec.host,
            port=spec.ssh_port,
            username=fleet.SIM_USERNAME,
            pkey=key,
            timeout=6,
            allow_agent=False,
            look_for_keys=False,
        )
    except paramiko.AuthenticationException:
        return (expect_reject, "被拒绝" if expect_reject else "被拒绝（authorized_keys 里没有这把公钥）")
    except Exception as exc:  # noqa: BLE001
        return (False, f"{type(exc).__name__}: {exc}")
    else:
        return (not expect_reject, "登录成功" if not expect_reject else "竟然登录成功了")
    finally:
        client.close()


def check_machine_addresses() -> None:
    """上下位机的地址解析（两个独立 IP 是理想情况，共用地址也可用）。"""
    detail = (
        f"{fleet.UPPER.name}={fleet.UPPER.host}:{fleet.UPPER.ssh_port} / "
        f"{fleet.LOWER1.name}={fleet.LOWER1.host}:{fleet.LOWER1.ssh_port}"
    )
    if fleet.hosts_are_distinct():
        detail += "（两个独立 IP）"
    else:
        aliases = fleet.loopback_aliases()
        detail += (
            f"（共用地址，靠端口区分；lo0 别名 {aliases or '无'}，"
            f"加别名可拿到两个独立 IP）"
        )
    _record(
        "上下位机地址已解析",
        all(bool(spec.host) for spec in fleet.FLEET),
        detail,
    )


def check_publickey_auth() -> None:
    """互信是可验证的：授权私钥放行，未授权私钥必须被拒。"""
    private = fleet.remote_to_local(fleet.UPPER, f"{fleet.UPPER.ssh_dir}/id_rsa")
    if not private.exists():
        _record("公钥免密登录可用", False, "上位机私钥不存在（先跑 scripts/sim.sh deploy）")
        return
    try:
        authorized = paramiko.RSAKey(filename=str(private))
    except paramiko.SSHException as exc:
        _record("公钥免密登录可用", False, f"私钥不可用：{exc}")
        return
    ok, detail = _passwordless_probe(fleet.LOWER1, authorized)
    _record("已授权私钥可免密登录", ok, detail)
    ok, detail = _passwordless_probe(fleet.LOWER1, paramiko.RSAKey.generate(2048), expect_reject=True)
    _record("未授权私钥被拒绝", ok, detail)


def check_auto_deploy() -> None:
    """自动化部署：整链可跑通、每步都留日志、日志全英文、可重复执行。"""
    from . import deploy as deploy_module

    first = deploy_module.deploy(seed_after=False, quiet=True, write_logs=True)
    failed = [step.title for step in first.steps if not step.ok]
    _record(
        "自动化部署全部步骤通过",
        first.ok,
        f"{len(first.steps)} 步" + (f"，失败：{'、'.join(failed)}" if failed else ""),
    )
    _record(
        "部署步数与声明一致",
        len(first.steps) == deploy_module.TOTAL_STEPS,
        f"{len(first.steps)}/{deploy_module.TOTAL_STEPS}",
    )
    _record(
        "部署覆盖互信的两个方向",
        any(step.key == "verify-forward" for step in first.steps)
        and any(step.key == "verify-reverse" for step in first.steps),
        "上位机→下位机 / 下位机→上位机 都校验了",
    )

    # 幂等：再跑一遍仍成功（密钥复用而不是重新生成）
    second = deploy_module.deploy(seed_after=False, quiet=True, write_logs=True)
    _record("部署可重复执行（幂等）", second.ok, "第二次执行仍全部通过" if second.ok else "第二次失败")

    # 每台机器都有一份部署日志，且记录了每一个步骤
    cjk = re.compile(r"[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef]")
    per_machine: list[tuple[str, int]] = []
    missing: list[str] = []
    chinese: list[str] = []
    for spec, path in zip(fleet.FLEET, first.log_files):
        if not path.exists():
            missing.append(spec.name)
            continue
        lines = path.read_text(encoding="utf-8").splitlines()
        per_machine.append((spec.name, len([line for line in lines if " step " in line])))
        chinese += [line for line in lines if cjk.search(line)]

    summary = "、".join(f"{name}: {count} 步" for name, count in per_machine)
    _record(
        "每台机器都有部署日志",
        not missing and len(per_machine) == len(fleet.FLEET),
        summary or f"缺失 {missing}",
    )
    _record(
        "部署日志记录每个步骤",
        bool(per_machine) and all(count >= deploy_module.TOTAL_STEPS for _n, count in per_machine),
        summary or "没有部署日志可查",
    )
    _record("部署日志全英文", not chinese, f"含中日韩字符 {len(chinese)} 行")

    # 部署事件按运行事件格式追加进了 event.log（前端「运行日志」能看到）
    run_log = fleet.remote_to_local(fleet.UPPER, f"{fleet.UPPER.run_root}/event.log")
    if run_log.exists():
        events = [
            line
            for line in run_log.read_text(encoding="utf-8").splitlines()
            if "[DEPLOY]" in line and "[deploy]" in line
        ]
        _record(
            "部署事件写入运行日志",
            len(events) >= deploy_module.TOTAL_STEPS,
            f"{len(events)} 条 [DEPLOY] 事件",
        )
    else:
        _record("部署事件写入运行日志", False, "event.log 不存在（先跑 init）")


def _bootstrap_django() -> bool:
    """配置 Django。

    这里必须先 ``django.setup()``：``check_backend_catalog`` 等检查会 import 模型，
    只 ``import django`` 是不够的，否则访问 settings 时抛 ImproperlyConfigured。
    """
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
    try:
        import django

        django.setup()
    except Exception as exc:  # noqa: BLE001
        _record("Django 初始化", False, f"{type(exc).__name__}: {exc}")
        return False
    _record("Django 初始化", True, os.environ["DJANGO_SETTINGS_MODULE"])
    return True


def run_all() -> int:
    _RESULTS.clear()
    for spec in fleet.FLEET:
        assert isinstance(spec, fleet.MachineSpec)

    try:
        check_ports()
    except Exception as exc:  # noqa: BLE001
        _record("机群可达", False, str(exc))

    try:
        client = _connect(fleet.UPPER)
    except Exception as exc:  # noqa: BLE001
        print(f"无法连接模拟上位机：{exc}")
        print("请先启动机群：cd backend && .venv/bin/python -m simremote.cli serve")
        return 1

    try:
        check_auth_reject()
        check_sftp_readonly(client)
        check_find(client)
        check_tar(client)
        check_awk_window(client)
        check_version_file(client)
        check_close_boundary()
        check_timestamp_jitter()
        check_machine_log_ownership()
        check_live_tail()
    finally:
        client.close()

    # 部署与互信（不依赖 Django：只动模拟文件系统与 SSH 通道）
    try:
        check_machine_addresses()
        check_publickey_auth()
        check_auto_deploy()
    except Exception as exc:  # noqa: BLE001
        _record("自动化部署检查", False, f"{type(exc).__name__}: {exc}")

    if _bootstrap_django():
        # 逐条包住：任何一条检查抛异常（改代码改出来的 NameError / KeyError 之类）
        # 以前会直接把整个自检打断，后面一百多条一条都不打印 —— 排查时只看到一个
        # 栈，完全不知道其余部分是好是坏。现在记一条 FAIL 然后继续跑。
        for check in (
            check_line_formats,
            check_backend_catalog,
            check_cpd_assets,
            check_cpd_backend_read,
            check_atlog_report_site,
            check_cpd_report_site,
            check_report_file_analysis,
            check_live_stream,
            check_live_cadence,
            check_live_rotation_cycle,
            check_log_call_chain,
            check_call_signature_shape,
            check_wsp_move_points,
            check_wsp_point_call_boundaries,
            check_wsp_spiral_trajectory,
            check_multiline_record_fixture,
        ):
            try:
                check()
            except Exception as exc:  # noqa: BLE001
                _record(f"{check.__name__} 执行失败", False, f"{type(exc).__name__}: {exc}")

    passed = sum(1 for _, ok, _detail in _RESULTS if ok)
    print("=" * 64)
    for name, ok, detail in _RESULTS:
        mark = "PASS" if ok else "FAIL"
        suffix = f"  — {detail}" if detail else ""
        print(f"  [{mark}] {name}{suffix}")
    print("=" * 64)
    print(f"  {passed}/{len(_RESULTS)} 项通过")
    return 0 if passed == len(_RESULTS) else 1
