"""Remote CPD workbooks: recursive discovery and row-time filtering, never mtime."""
from __future__ import annotations

import json
import posixpath
import sqlite3
import stat
import tempfile
from contextlib import contextmanager
from datetime import date, datetime, time
from pathlib import PurePosixPath
from zipfile import ZipFile

from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from apps.common.services.ssh import ssh_session
from apps.reports.services import _safe_part

TIME_COLUMNS = ('timestamp', 'time', 'date', '时间', '时间戳')


def data_root(environment, subsystem: str, module: str) -> str:
    username = _safe_part(environment.upper_machine.username, '上位机用户名')
    return f'/data/{username}/cpd_data/{_safe_part(subsystem, "子系统").lower()}/{_safe_part(module, "模块").lower()}'


def safe_file(root: str, path: str) -> str:
    normalized = posixpath.normpath(path)
    if not normalized.startswith(root + '/') or not normalized.lower().endswith('.xlsx') or '..' in PurePosixPath(path).parts:
        raise ValueError('只能读取当前环境测校数据目录内的 .xlsx 文件')
    return normalized


def parse_time(value, *, epoch=None, anchor: datetime | None = None):
    if value is None or value == '':
        return None
    try:
        if isinstance(value, bool): return None
        if isinstance(value, (int, float)):
            value = from_excel(value, epoch=epoch) if epoch else None
        if isinstance(value, time):
            value = datetime.combine(anchor.date(), value) if anchor else None
        if isinstance(value, date) and not isinstance(value, datetime):
            value = datetime.combine(value, time())
        if isinstance(value, str):
            value = datetime.fromisoformat(value.strip().replace('Z', '+00:00').replace('/', '-'))
        if not isinstance(value, datetime): return None
        return timezone.make_aware(value, timezone.get_current_timezone()) if timezone.is_naive(value) else value
    except (ValueError, TypeError, OverflowError):
        return None


def time_window(start_time: str, end_time: str):
    start, end = parse_time(start_time), parse_time(end_time)
    if not start or not end or start > end:
        raise ValueError('必须提供有效的测校开始/结束时间')
    return start, end


def walk_xlsx(sftp, root: str):
    # Iterative traversal has no depth limit. Do not follow links outside the root.
    pending = [root]
    visited = set()
    while pending:
        path = pending.pop()
        if path in visited: continue
        visited.add(path)
        for item in sftp.listdir_attr(path):
            child = posixpath.join(path, item.filename)
            if item.filename in {'.', '..'} or '/' in item.filename: continue
            if stat.S_ISDIR(item.st_mode): pending.append(child)
            elif stat.S_ISREG(item.st_mode) and item.filename.lower().endswith('.xlsx') and not item.filename.startswith('~$'):
                yield {'name': item.filename, 'path': child, 'size': item.st_size}


@contextmanager
def remote_workbook(sftp, root, path):
    path = safe_file(root, path)
    # Reject symlink escapes on direct preview requests too.
    if not sftp.normalize(path).startswith(sftp.normalize(root).rstrip('/') + '/'):
        raise ValueError('文件链接超出测校目录')
    if sftp.stat(path).st_size > 128 * 1024 * 1024:
        raise ValueError('Excel 超过 128 MiB，需拆分工作簿后预览')
    with tempfile.SpooledTemporaryFile(max_size=4 * 1024 * 1024) as tmp:
        with sftp.open(path, 'rb') as remote:
            total = 0
            while chunk := remote.read(1024 * 1024):
                total += len(chunk)
                if total > 128 * 1024 * 1024: raise ValueError('Excel 超过读取预算')
                tmp.write(chunk)
        tmp.seek(0)
        with ZipFile(tmp) as archive:
            if sum(item.file_size for item in archive.infolist()) > 512 * 1024 * 1024:
                raise ValueError('Excel 解压内容超过 512 MiB')
        tmp.seek(0)
        workbook = load_workbook(tmp, read_only=True, data_only=True)
        try: yield workbook
        finally: workbook.close()


def sheet_rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    for number, row in enumerate(iterator, 1):
        if number > 20: break
        headers = [str(cell or '').strip() for cell in row]
        lowered = [h.casefold() for h in headers]
        column = next((lowered.index(name) for name in TIME_COLUMNS if name in lowered), None)
        if column is not None:
            return headers, column, enumerate(iterator, number + 1)
    return [], None, iter(())


def inspect_sheet(sheet, epoch, start, end):
    headers, column, rows = sheet_rows(sheet)
    if column is None:
        return {'name': sheet.title, 'time_column': None, 'matched_rows': 0, 'invalid_time_rows': 0, 'filter_status': 'missing_time_column'}
    count = invalid = 0
    anchor = start if start.date() == end.date() else None
    for _, row in rows:
        ts = parse_time(row[column] if column < len(row) else None, epoch=epoch, anchor=anchor)
        if ts is None: invalid += 1
        elif start <= ts <= end: count += 1
    return {'name': sheet.title, 'time_column': headers[column], 'matched_rows': count, 'invalid_time_rows': invalid, 'filter_status': 'filtered'}


def find_cpd_excel_files(environment, subsystem: str, module: str, start_time: str, end_time: str):
    start, end = time_window(start_time, end_time)
    root = data_root(environment, subsystem, module)
    files, errors = [], []
    with ssh_session(environment.upper_machine) as lease:
        sftp = lease.client.open_sftp()
        try:
            for item in walk_xlsx(sftp, root):
                try:
                    with remote_workbook(sftp, root, item['path']) as workbook:
                        sheets = [inspect_sheet(sheet, workbook.epoch, start, end) for sheet in workbook.worksheets]
                    if any(s['matched_rows'] or s['filter_status'] == 'missing_time_column' for s in sheets):
                        files.append({**item, 'sheets': sheets})
                except Exception as exc:
                    errors.append({'path': item['path'], 'error': str(exc)})
        finally: sftp.close()
    return {'root': root, 'files': sorted(files, key=lambda f: f['path']), 'errors': errors, 'filter_basis': 'excel_time_column'}


def preview_sheet(workbook, sheet_name, start, end, *, page=1, page_size=100, sort_column=None, descending=False):
    if sheet_name not in workbook.sheetnames: raise ValueError('Sheet 不存在')
    page, page_size = max(1, int(page)), max(1, min(500, int(page_size)))
    headers, column, rows = sheet_rows(workbook[sheet_name])
    if column is None:
        return {'sheet': sheet_name, 'headers': [], 'rows': [], 'total': 0, 'page': page, 'page_size': page_size, 'filter_status': 'missing_time_column', 'warning': '未找到时间列，不能确定哪些行属于此次测校；没有用文件修改时间代替'}
    sort_index = int(sort_column) if sort_column is not None and str(sort_column) != '' else None
    if sort_index is not None and not 0 <= sort_index < len(headers): raise ValueError('排序列无效')
    def scalar(v):
        return v.isoformat(sep=' ') if isinstance(v, datetime) else v.isoformat() if isinstance(v, (date, time)) else v if isinstance(v, (str, int, float, bool)) or v is None else str(v)
    invalid = 0
    # SQLite temp storage keeps sorting/pagination bounded in memory even for large sheets.
    with tempfile.TemporaryDirectory(prefix='tracelens-cpd-') as folder:
        connection = sqlite3.connect(folder + '/rows.sqlite3')
        try:
            connection.execute('CREATE TABLE rows (line INTEGER, value TEXT, sort_type INTEGER, sort_number REAL, sort_text TEXT)')
            for number, row in rows:
                ts = parse_time(row[column] if column < len(row) else None, epoch=workbook.epoch, anchor=start if start.date() == end.date() else None)
                if ts is None: invalid += 1; continue
                if not start <= ts <= end: continue
                values = [scalar(v) for v in row]
                value = values[sort_index] if sort_index is not None and sort_index < len(values) else number
                numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                connection.execute('INSERT INTO rows VALUES (?,?,?,?,?)', (number, json.dumps(values, ensure_ascii=False), 0 if numeric else 1, value if numeric else 0, str(value or '')))
            total = connection.execute('SELECT COUNT(*) FROM rows').fetchone()[0]
            direction = 'DESC' if descending else 'ASC'
            ordering = f'sort_type, sort_number {direction}, sort_text {direction}, line' if sort_index is not None else 'line'
            records = connection.execute(f'SELECT line,value FROM rows ORDER BY {ordering} LIMIT ? OFFSET ?', (page_size, (page - 1) * page_size)).fetchall()
        finally: connection.close()
    return {'sheet': sheet_name, 'headers': headers, 'rows': [{'line': line, 'cells': json.loads(value)} for line, value in records], 'total': total, 'page': page, 'page_size': page_size, 'time_column': headers[column], 'invalid_time_rows': invalid, 'filter_status': 'filtered'}


def preview_cpd_excel(environment, subsystem, module, path, sheet, start_time, end_time, **options):
    start, end = time_window(start_time, end_time)
    root = data_root(environment, subsystem, module)
    with ssh_session(environment.upper_machine) as lease:
        sftp = lease.client.open_sftp()
        try:
            with remote_workbook(sftp, root, path) as workbook:
                return {'path': path, 'sheets': workbook.sheetnames, **preview_sheet(workbook, sheet or workbook.sheetnames[0], start, end, **options)}
        finally: sftp.close()
